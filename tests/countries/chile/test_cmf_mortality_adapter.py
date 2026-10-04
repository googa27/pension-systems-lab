from __future__ import annotations

import hashlib
from collections.abc import Callable
from dataclasses import replace
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import chile_demographic_pde.countries.chile.pensions as pensions_module
from chile_demographic_pde.core.errors import (
    DataContractError,
    MissingOfficialDataError,
)
from chile_demographic_pde.countries.chile.official_finance_registry import (
    ReviewedOfficialAsset,
    ReviewedOfficialReleaseEntry,
    ReviewedOfficialReleaseRegistry,
)
from chile_demographic_pde.countries.chile.pensions import (
    CMF_TM2020_HISTORICAL_CONTRACT_V1,
    CmfMortalityRelease,
    CmfMortalityTableAdapter,
    CmfMortalityWorkbookBlock,
    CmfMortalityWorkbookContract,
    parse_cmf_tm2020_historical,
)
from chile_demographic_pde.data.acquisition import AcquiredAsset
from chile_demographic_pde.data.domain_schemas import (
    validate_regulatory_mortality,
)
from chile_demographic_pde.data.identity import IdentityRegistry
from chile_demographic_pde.data.releases import (
    ReleaseManifest,
    RequiredReleaseAsset,
    verify_release_assets,
)
from chile_demographic_pde.data.roles import ReleaseMissingnessReason

_SOURCE_KEY = "cmf_tm2020_historical_xlsx"
_REGISTRY_KEY = "cmf_tm2020_historical"
_REGISTRY_VERSION = "chile-reviewed-official-finance-v1"
_RELEASE_SCHEMA_VERSION = "cmf-tm2020-historical-xlsx-v1"
_CONTRACT_ID = "cmf-tm2020-historical-workbook-v1"
_RULE_ID = "NCG_495_SP_NCG_306"
_RULE_RELEASED_AT = date(2023, 2, 24)
_EFFECTIVE_START = date(2023, 7, 1)
_MAXIMUM_EFFECTIVE_END = date(2029, 7, 1)
_AVAILABLE_AT = datetime(2026, 7, 19, 16, tzinfo=UTC)
_MISSING_RELEASE_DATE = ReleaseMissingnessReason.PUBLISHER_DATE_NOT_AVAILABLE
_NOTE = "*: A partir del año 2036, los factores de mejoramiento se mantienen constantes"


def _qx(age: int) -> float:
    """Return a deterministic parser-fixture probability, never production data."""

    return (age + 1) / 1_000


def _improvement(age: int, year: int) -> float:
    """Return a deterministic parser-fixture factor, never production data."""

    return (year - 2020) / 10_000 + age / 1_000_000


def _write_reviewed_block(
    worksheet: object,
    *,
    start_column: int,
    title: str,
    age_start: int,
    note_row: int,
) -> None:
    # ``worksheet`` is deliberately kept local to the test-fixture builder;
    # production adapters never see generated observations as official data.
    ws = worksheet
    ws.cell(row=1, column=start_column, value=title)  # type: ignore[attr-defined]
    ws.cell(row=2, column=start_column, value="Edad")  # type: ignore[attr-defined]
    ws.cell(row=2, column=start_column + 1, value="qx 2020")  # type: ignore[attr-defined]
    ws.cell(  # type: ignore[attr-defined]
        row=2,
        column=start_column + 2,
        value="Factores de mejoramiento AAx,t",
    )
    for offset, year in enumerate(range(2021, 2037), start=2):
        ws.cell(  # type: ignore[attr-defined]
            row=3,
            column=start_column + offset,
            value=year if year < 2036 else "2036*",
        )
    for row, age in enumerate(range(age_start, 111), start=4):
        ws.cell(row=row, column=start_column, value=age)  # type: ignore[attr-defined]
        ws.cell(row=row, column=start_column + 1, value=_qx(age))  # type: ignore[attr-defined]
        for offset, year in enumerate(range(2021, 2037), start=2):
            ws.cell(  # type: ignore[attr-defined]
                row=row,
                column=start_column + offset,
                value=_improvement(age, year),
            )
    ws.cell(row=note_row, column=start_column, value=_NOTE)  # type: ignore[attr-defined]


def _generated_historical_workbook_bytes() -> bytes:
    """Build the reviewed geometry in memory for parser tests only."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "chile-demographic-pde test fixture"
    workbook.properties.created = datetime(2026, 7, 19)
    workbook.properties.modified = datetime(2026, 7, 19)

    layouts = (
        (
            "Vejez-Mujeres",
            13,
            "TABLA RV-2020-MUJERES",
            20,
            95,
            (
                (1, "TABLA RV-2004-MUJERES"),
                (5, "TABLA RV-2009-MUJERES"),
                (9, "TABLA RV-2014-MUJERES"),
            ),
        ),
        (
            "Invalidez-Mujeres",
            9,
            "TABLA MI-2020-MUJERES",
            0,
            115,
            (
                (1, "TABLA MI-2006-MUJERES"),
                (5, "TABLA MI-2014 - MUJERES"),
            ),
        ),
        (
            "Sobrevivencia-Mujeres",
            9,
            "TABLA B-2020-MUJERES",
            0,
            115,
            (
                (1, "TABLA B-2006-MUJERES"),
                (5, "TABLA B-2014-MUJERES"),
            ),
        ),
        (
            "Vejez-Hombres",
            13,
            "TABLA CB-2020-HOMBRES",
            0,
            115,
            (
                (1, "TABLA RV-2004-HOMBRES"),
                (5, "TABLA RV-2009-HOMBRES"),
                (9, "TABLA CB-2014-HOMBRES"),
            ),
        ),
        (
            "Invalidez-Hombres",
            9,
            "TABLA MI-2020 - HOMBRES",
            0,
            115,
            (
                (1, "TABLA MI-2006-HOMBRES"),
                (5, "TABLA MI-2014 - HOMBRES"),
            ),
        ),
        (
            "Sobrevivencia-Hombres",
            9,
            "TABLA CB-2020-HOMBRES",
            0,
            115,
            (
                (1, "TABLA B-2006-HOMBRES"),
                (5, "TABLA CB-2014-HOMBRES"),
            ),
        ),
    )
    for (
        sheet_name,
        start_column,
        title,
        age_start,
        note_row,
        earlier_titles,
    ) in layouts:
        worksheet = workbook.create_sheet(sheet_name)
        for column, earlier_title in earlier_titles:
            worksheet.cell(row=1, column=column, value=earlier_title)
        _write_reviewed_block(
            worksheet,
            start_column=start_column,
            title=title,
            age_start=age_start,
            note_row=note_row,
        )

    destination = BytesIO()
    workbook.save(destination)
    return destination.getvalue()


def _mutated_workbook(raw: bytes, mutation: str) -> bytes:
    workbook = load_workbook(BytesIO(raw), data_only=False)
    if mutation == "sheet":
        workbook["Vejez-Mujeres"].title = "Vejez Mujeres"
    elif mutation == "dimension":
        workbook["Vejez-Mujeres"]["AE1"] = "unexpected"
    elif mutation == "title":
        workbook["Invalidez-Hombres"]["I1"] = "TABLA MI-2020-HOMBRES"
    elif mutation == "header":
        workbook["Sobrevivencia-Mujeres"]["J2"] = "qx"
    elif mutation == "year":
        workbook["Vejez-Hombres"]["O3"] = 2020
    elif mutation == "note":
        workbook["Invalidez-Mujeres"]["I115"] = "changed note"
    elif mutation == "age_support":
        workbook["Sobrevivencia-Mujeres"]["I14"] = 999
    elif mutation == "qx":
        workbook["Vejez-Mujeres"]["N4"] = -0.1
    elif mutation == "improvement":
        workbook["Invalidez-Mujeres"]["K4"] = -0.01
    elif mutation == "duplicate_cb":
        workbook["Sobrevivencia-Hombres"]["J69"] = 0.5
    elif mutation == "duplicate_cb_nonemitted_cell":
        workbook["Sobrevivencia-Hombres"]["L2"] = "unexpected duplicate-only value"
    else:
        raise AssertionError(f"unknown workbook mutation {mutation!r}")
    destination = BytesIO()
    workbook.save(destination)
    return destination.getvalue()


def _test_registry(raw: bytes) -> ReviewedOfficialReleaseRegistry:
    digest = hashlib.sha256(raw).hexdigest()
    entry = ReviewedOfficialReleaseEntry(
        registry_version=_REGISTRY_VERSION,
        registry_key=_REGISTRY_KEY,
        release_schema_version=_RELEASE_SCHEMA_VERSION,
        primary_fact_source_key=_SOURCE_KEY,
        assets=(
            ReviewedOfficialAsset(
                source_key=_SOURCE_KEY,
                sha256=digest,
                released_at=None,
                release_missingness_reason=_MISSING_RELEASE_DATE,
            ),
        ),
        contract_id=_CONTRACT_ID,
        review_id="generated-parser-fixture-only",
    )
    return ReviewedOfficialReleaseRegistry.create(
        registry_version=_REGISTRY_VERSION,
        entries=(entry,),
    )


def _release(
    tmp_path: Path,
    raw: bytes,
    identity_registry: IdentityRegistry,
    *,
    available_at: datetime = _AVAILABLE_AT,
    vintage: str = "TM-2020-generated-parser-fixture",
) -> CmfMortalityRelease:
    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "tm2020-generated-parser-fixture.xlsx"
    path.write_bytes(raw)
    digest = hashlib.sha256(raw).hexdigest()
    acquired = AcquiredAsset(
        key=_SOURCE_KEY,
        path=path,
        sha256=digest,
        retrieved_at=available_at,
    )
    manifest = ReleaseManifest.create(
        schema_version=_RELEASE_SCHEMA_VERSION,
        assets=(
            RequiredReleaseAsset(
                source_key=_SOURCE_KEY,
                sha256=digest,
                available_at=available_at,
                released_at=None,
                release_missingness_reason=_MISSING_RELEASE_DATE,
            ),
        ),
        primary_fact_source_key=_SOURCE_KEY,
        identity_registry=identity_registry,
    )
    return CmfMortalityRelease(
        assets=verify_release_assets(manifest, (acquired,)),
        vintage=vintage,
        rule_id=_RULE_ID,
        rule_released_at=_RULE_RELEASED_AT,
        effective_start=_EFFECTIVE_START,
        effective_end=None,
        maximum_effective_end=_MAXIMUM_EFFECTIVE_END,
    )


@pytest.fixture(scope="module")
def generated_mortality_workbook_bytes() -> bytes:
    return _generated_historical_workbook_bytes()


@pytest.fixture(scope="module")
def mortality_workbook_contract() -> CmfMortalityWorkbookContract:
    return CMF_TM2020_HISTORICAL_CONTRACT_V1


@pytest.fixture
def test_official_registry(
    generated_mortality_workbook_bytes: bytes,
) -> ReviewedOfficialReleaseRegistry:
    return _test_registry(generated_mortality_workbook_bytes)


@pytest.fixture
def cmf_mortality_release(
    tmp_path: Path,
    generated_mortality_workbook_bytes: bytes,
) -> CmfMortalityRelease:
    return _release(
        tmp_path,
        generated_mortality_workbook_bytes,
        IdentityRegistry(),
    )


def test_historical_contract_is_a_closed_six_block_value() -> None:
    contract = CMF_TM2020_HISTORICAL_CONTRACT_V1

    assert isinstance(contract, CmfMortalityWorkbookContract)
    assert len(contract.blocks) == 6
    assert all(isinstance(block, CmfMortalityWorkbookBlock) for block in contract.blocks)


def test_pure_parser_emits_exact_unique_fact_cardinality(
    generated_mortality_workbook_bytes: bytes,
    mortality_workbook_contract: CmfMortalityWorkbookContract,
) -> None:
    facts = parse_cmf_tm2020_historical(
        generated_mortality_workbook_bytes,
        mortality_workbook_contract,
    )
    base = [fact for fact in facts if fact.variable == "regulatory_mortality_probability"]
    improvements = [
        fact for fact in facts if fact.variable == "regulatory_mortality_improvement_factor"
    ]

    assert isinstance(facts, tuple)
    assert len(facts) == 9_095
    assert len(base) == 535
    assert len(improvements) == 8_560
    assert {fact.table_id for fact in facts} == {
        "RV-M-2020",
        "MI-M-2020",
        "B-M-2020",
        "CB-H-2020",
        "MI-H-2020",
    }
    assert sum(fact.table_id == "CB-H-2020" for fact in base) == 111


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("sheet", r"sheet|workbook"),
        ("dimension", r"dimension|worksheet"),
        ("title", r"title|block"),
        ("header", r"header"),
        ("year", r"year|header"),
        ("note", r"note"),
        ("age_support", r"age"),
        ("qx", r"qx|probability"),
        ("improvement", r"improvement|factor"),
        ("duplicate_cb", r"duplicate|CB-H"),
        ("duplicate_cb_nonemitted_cell", r"duplicate|CB-H"),
    ],
)
def test_pure_parser_fails_closed_on_every_reviewed_workbook_atom(
    mutation: str,
    message: str,
    generated_mortality_workbook_bytes: bytes,
    mortality_workbook_contract: CmfMortalityWorkbookContract,
) -> None:
    changed = _mutated_workbook(
        generated_mortality_workbook_bytes,
        mutation,
    )

    with pytest.raises(DataContractError, match=message):
        parse_cmf_tm2020_historical(changed, mortality_workbook_contract)


def test_adapter_requires_injected_authorization_for_generated_parser_fixture(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    with pytest.raises(MissingOfficialDataError):
        CmfMortalityTableAdapter().normalize_tables(
            cmf_mortality_release,
            identity_registry=IdentityRegistry(),
        )

    batch = CmfMortalityTableAdapter(test_official_registry).normalize_tables(
        cmf_mortality_release,
        identity_registry=IdentityRegistry(),
    )
    assert len(batch.observations) == 9_095


def test_authorization_failure_precedes_any_workbook_open(
    cmf_mortality_release: CmfMortalityRelease,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def forbidden_open(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("unreviewed bytes were opened")

    monkeypatch.setattr(type(cmf_mortality_release.assets), "open", forbidden_open)

    with pytest.raises(MissingOfficialDataError):
        CmfMortalityTableAdapter().normalize_tables(
            cmf_mortality_release,
            identity_registry=IdentityRegistry(),
        )


def test_adapter_keeps_base_qx_and_improvement_factors_separate(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    observations = (
        CmfMortalityTableAdapter(test_official_registry)
        .normalize_tables(
            cmf_mortality_release,
            identity_registry=IdentityRegistry(),
        )
        .observations
    )

    assert observations["variable"].value_counts().to_dict() == {
        "regulatory_mortality_improvement_factor": 8_560,
        "regulatory_mortality_probability": 535,
    }
    assert "projected_qx" not in set(observations["variable"])
    base = observations["variable"].eq("regulatory_mortality_probability")
    improvement = observations["variable"].eq("regulatory_mortality_improvement_factor")
    assert set(observations.loc[base, "semantic_kind"]) == {"probability"}
    assert set(observations.loc[base, "population_basis"]) == {"pensioner"}
    assert observations.loc[base, "improvement_year"].isna().all()
    assert set(observations.loc[improvement, "semantic_kind"]) == {"index"}
    assert set(observations.loc[improvement, "population_basis"]) == {"not_applicable"}
    assert set(observations.loc[improvement, "improvement_year"]) == set(range(2021, 2037))
    assert observations["transformation_id"].isna().all()
    assert observations["transformation_version"].isna().all()


def test_adapter_preserves_table_sex_age_and_value_mappings(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    observations = (
        CmfMortalityTableAdapter(test_official_registry)
        .normalize_tables(
            cmf_mortality_release,
            identity_registry=IdentityRegistry(),
        )
        .observations
    )
    base = observations.loc[observations["variable"].eq("regulatory_mortality_probability")]

    assert base.groupby("mortality_table_id")["sex"].unique().map(tuple).to_dict() == {
        "B-M-2020": ("female",),
        "CB-H-2020": ("male",),
        "MI-H-2020": ("male",),
        "MI-M-2020": ("female",),
        "RV-M-2020": ("female",),
    }
    assert base.groupby("mortality_table_id")["age"].agg(["min", "max"]).to_dict(
        orient="index"
    ) == {
        "B-M-2020": {"min": 0, "max": 110},
        "CB-H-2020": {"min": 0, "max": 110},
        "MI-H-2020": {"min": 0, "max": 110},
        "MI-M-2020": {"min": 0, "max": 110},
        "RV-M-2020": {"min": 20, "max": 110},
    }
    anchor = base.loc[base["mortality_table_id"].eq("CB-H-2020") & base["age"].eq(65)].squeeze()
    assert anchor["value"] == pytest.approx(_qx(65))
    assert set(observations["population_segment"]) == {"pensioner"}


def test_reference_rule_and_effective_dates_remain_distinct(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    observations = (
        CmfMortalityTableAdapter(test_official_registry)
        .normalize_tables(
            cmf_mortality_release,
            identity_registry=IdentityRegistry(),
        )
        .observations
    )

    assert cmf_mortality_release.rule_id == _RULE_ID
    assert cmf_mortality_release.rule_released_at == _RULE_RELEASED_AT
    assert set(observations["reference_year"]) == {2020}
    assert set(observations["effective_start"]) == {_EFFECTIVE_START}
    assert observations["effective_end"].isna().all()
    assert set(observations["maximum_effective_end"]) == {_MAXIMUM_EFFECTIVE_END}
    assert set(observations["regulatory_basis"]) == {"TM_2020"}


def test_normalized_batch_satisfies_strict_regulatory_mortality_schema(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    batch = CmfMortalityTableAdapter(test_official_registry).normalize_tables(
        cmf_mortality_release,
        identity_registry=IdentityRegistry(),
    )

    validated = validate_regulatory_mortality(batch.observations)
    pd.testing.assert_frame_equal(validated, batch.observations)
    assert len(batch.identities.records) == 18_191


def test_failed_domain_validation_does_not_mutate_caller_identity_registry(
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    identity_registry = IdentityRegistry()
    before = identity_registry.snapshot()

    def fail_domain_validation(_frame: object) -> object:
        raise DataContractError("forced mortality-domain validation failure")

    monkeypatch.setattr(
        pensions_module,
        "validate_regulatory_mortality",
        fail_domain_validation,
    )

    with pytest.raises(DataContractError, match="forced"):
        CmfMortalityTableAdapter(test_official_registry).normalize_tables(
            cmf_mortality_release,
            identity_registry=identity_registry,
        )

    assert identity_registry.snapshot() == before


def test_fact_identity_is_stable_across_release_revisions(
    tmp_path: Path,
    generated_mortality_workbook_bytes: bytes,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    first_registry = IdentityRegistry()
    second_registry = IdentityRegistry()
    first_release = _release(
        tmp_path / "first",
        generated_mortality_workbook_bytes,
        first_registry,
    )
    second_release = _release(
        tmp_path / "second",
        generated_mortality_workbook_bytes,
        second_registry,
        available_at=_AVAILABLE_AT.replace(day=20),
        vintage="TM-2020-generated-parser-fixture-retrieved-later",
    )
    adapter = CmfMortalityTableAdapter(test_official_registry)

    first = adapter.normalize_tables(
        first_release,
        identity_registry=first_registry,
    ).observations.sort_values(
        ["mortality_table_id", "age", "variable", "improvement_year"],
        na_position="first",
    )
    second = adapter.normalize_tables(
        second_release,
        identity_registry=second_registry,
    ).observations.sort_values(
        ["mortality_table_id", "age", "variable", "improvement_year"],
        na_position="first",
    )

    assert list(first["fact_id"]) == list(second["fact_id"])
    assert set(first["release_id"]) == set(second["release_id"])
    assert set(first["observation_id"]).isdisjoint(set(second["observation_id"]))


@pytest.mark.parametrize(
    "change",
    [
        lambda release: replace(release, rule_id="caller_rule"),
        lambda release: replace(
            release,
            rule_released_at=date(2023, 2, 25),
        ),
        lambda release: replace(
            release,
            effective_start=date(2023, 7, 2),
        ),
        lambda release: replace(
            release,
            effective_end=date(2025, 1, 1),
        ),
        lambda release: replace(
            release,
            maximum_effective_end=date(2029, 7, 2),
        ),
    ],
)
def test_adapter_rejects_changed_regulatory_release_semantics(
    change: Callable[[CmfMortalityRelease], CmfMortalityRelease],
    cmf_mortality_release: CmfMortalityRelease,
    test_official_registry: ReviewedOfficialReleaseRegistry,
) -> None:
    with pytest.raises(DataContractError, match=r"rule|effective|TM-2020"):
        CmfMortalityTableAdapter(test_official_registry).normalize_tables(
            change(cmf_mortality_release),
            identity_registry=IdentityRegistry(),
        )
