from __future__ import annotations

import hashlib
import json
from collections.abc import Callable, Mapping, Sequence
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path
from random import Random
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from chile_demographic_pde.core.errors import (
    DataContractError,
    MissingOfficialDataError,
)
from chile_demographic_pde.countries.chile import pensions as pensions_module
from chile_demographic_pde.countries.chile.official_finance_registry import (
    ReviewedOfficialAsset,
    ReviewedOfficialReleaseEntry,
    ReviewedOfficialReleaseRegistry,
)
from chile_demographic_pde.countries.chile.pensions import (
    CmfMortalityConstructionContract,
    CmfMortalityRelease,
    CmfMortalityTableAdapter,
    parse_cmf_tm2020_construction,
)
from chile_demographic_pde.data.acquisition import AcquiredAsset
from chile_demographic_pde.data.identity import IdentityRegistry
from chile_demographic_pde.data.releases import (
    ReleaseManifest,
    RequiredReleaseAsset,
    VerifiedReleaseAssets,
    verify_release_assets,
)
from chile_demographic_pde.data.roles import ReleaseMissingnessReason

_MICRODATA_MEMBER = "bd/bd_publica.txt"
_CONTROL_MEMBERS = {
    "RV-M-2020": (
        "tablas_mortalidad/rv 2020-mujeres.xlsx",
        "RV-2020, Mujeres",
    ),
    "CB-H-2020": (
        "tablas_mortalidad/cb 2020-hombres.xlsx",
        "CB-2020, Hombres",
    ),
    "B-M-2020": (
        "tablas_mortalidad/b 2020-mujeres.xlsx",
        "B-2020, Mujeres",
    ),
    "MI-M-2020": (
        "tablas_mortalidad/mi 2020-mujeres.xlsx",
        "MI-2020, Mujeres",
    ),
    "MI-H-2020": (
        "tablas_mortalidad/mi 2020-hombres.xlsx",
        "MI-2020, Hombres",
    ),
}
_INE_MEMBERS = (
    "tablas_ine/tablas_ine-hombres.xlsx",
    "tablas_ine/tablas_ine-mujeres.xlsx",
)
_EXPECTED_MEMBERS = tuple(
    sorted((_MICRODATA_MEMBER, *_INE_MEMBERS, *(value[0] for value in _CONTROL_MEMBERS.values())))
)
_TABLE_ORDER = ("RV-M-2020", "CB-H-2020", "B-M-2020", "MI-M-2020", "MI-H-2020")
_DESCRIPTOR_KEY = "cmf_tm2020_descriptor_pdf"
_DEVIATION_KEY = "cmf_tm2020_public_original_deviation_v1"
_ZIP_KEY = "cmf_tm2020_construction_zip"
_DESCRIPTOR_RAW = b"generated descriptor placeholder; not an official observation"
_DESCRIPTOR_SHA256 = hashlib.sha256(_DESCRIPTOR_RAW).hexdigest()
_REGISTRY_VERSION = "chile-reviewed-official-finance-v1"
_REGISTRY_KEY = "cmf_tm2020_construction_diagnostic"
_RELEASE_SCHEMA = "cmf-tm2020-construction-diagnostic-v1"
_CONTRACT_ID = "cmf-tm2020-construction-diagnostic-v1"
_MISSING_DATE = ReleaseMissingnessReason.PUBLISHER_DATE_NOT_AVAILABLE
_NOW = datetime(2026, 7, 19, 12, tzinfo=UTC)

_MICRODATA_COLUMNS = (
    "ORIGEN",
    "ID",
    "TABLA",
    "FEC_VIG",
    "FIN_EXPOSIC",
    "SEXO",
    "COD_REL",
    "BEN_COD_INV",
    "FEC_NAC",
    "FEC_FALL",
)
_AUDIT_COLUMNS = (
    "table_id",
    "age",
    "public_exposed",
    "public_deaths",
    "public_qx",
    "workbook_exposed",
    "workbook_deaths",
    "workbook_qx",
    "published_delta_exposed",
    "published_delta_deaths",
    "exposed_relation_ok",
    "deaths_relation_ok",
    "qx_reconciled",
)

Count = tuple[int, int]
ControlCounts = Mapping[str, Mapping[int, Count]]


def _row(
    table: int,
    sex: str,
    invalidity: str,
    *,
    relationship: int = 99,
    pension_date: int = 20200101,
    end_exposure: int = 99991231,
    birth_date: int = 19800101,
    death_date: int = 0,
    token: str = "opaque-fixture-token",
) -> dict[str, object]:
    return {
        "ORIGEN": 1,
        "ID": token,
        "TABLA": table,
        "FEC_VIG": pension_date,
        "FIN_EXPOSIC": end_exposure,
        "SEXO": sex,
        "COD_REL": relationship,
        "BEN_COD_INV": invalidity,
        "FEC_NAC": birth_date,
        "FEC_FALL": death_date,
    }


def _closed_routing_rows() -> list[dict[str, object]]:
    """Cover every reviewed route while contributing no anniversary exposure."""

    return [
        _row(1, "F", "N", relationship=10, token="route-b"),
        _row(2, "F", "N", token="route-rv"),
        _row(3, "M", "N", token="route-cb"),
        _row(4, "F", "P", token="route" + "-mi-f-p"),
        _row(4, "F", "T", relationship=10, token="route" + "-mi-f-t"),
        _row(4, "M", "P", token="route" + "-mi-m-p"),
        _row(4, "M", "T", relationship=10, token="route" + "-mi-m-t"),
    ]


def _microdata_bytes(rows: Sequence[Mapping[str, object]]) -> bytes:
    lines = [";".join(_MICRODATA_COLUMNS)]
    lines.extend(";".join(str(row[column]) for column in _MICRODATA_COLUMNS) for row in rows)
    return ("\n".join(lines) + "\n").encode()


def _mortality_workbook(
    table_id: str,
    counts: Mapping[int, Count],
) -> bytes:
    workbook = Workbook()
    raw = workbook.active
    raw.title = "qx_brutos"
    raw["B2"] = "GENERATED TEST CONTROL"
    headers = ("Edad", "Expuestos", "Muertos", "qx_bruto")
    for column, header in enumerate(headers, start=2):
        raw.cell(row=6, column=column, value=header)
    for age in range(111):
        exposed, deaths = counts.get(age, (0, 0))
        raw.cell(row=age + 7, column=2, value=age)
        raw.cell(row=age + 7, column=3, value=exposed)
        raw.cell(row=age + 7, column=4, value=deaths)
        raw.cell(
            row=age + 7,
            column=5,
            value=(deaths / exposed if exposed else None),
        )
    # The reviewed source has a styled, blank K117 cell in its exact B2:K117
    # dimension. Creating the blank cell preserves that bounded shape.
    raw.cell(row=117, column=11).fill = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF",
    )

    reviewed_name = _CONTROL_MEMBERS[table_id][1]
    reviewed = workbook.create_sheet(reviewed_name)
    reviewed.cell(row=2, column=1, value=f"Generated {table_id} parser fixture")
    reviewed.cell(row=7, column=1, value="Edad")
    reviewed.cell(row=7, column=2, value="Tasas de mortalidad qx")
    reviewed.cell(row=7, column=3, value="Factores de mejoramiento AAx,t")
    for offset, year in enumerate(range(2016, 2037), start=3):
        reviewed.cell(
            row=8,
            column=offset,
            value=year if year < 2036 else "2036*",
        )
    age_start = 20 if table_id == "RV-M-2020" else 0
    for offset, age in enumerate(range(age_start, 111), start=9):
        reviewed.cell(row=offset, column=1, value=age)
        reviewed.cell(row=offset, column=2, value=0.0)
        for column in range(3, 24):
            reviewed.cell(row=offset, column=column, value=0.0)
    note_row = 104 if table_id == "RV-M-2020" else 124
    reviewed.cell(
        row=note_row,
        column=1,
        value=("*: A partir del año 2036, los factores de mejoramiento se mantienen constantes"),
    )
    reviewed.cell(row=note_row, column=23)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _ine_workbook(sex: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "TM_históricas"
    sheet.cell(
        row=2,
        column=1,
        value=(
            "Tablas de mortalidad históricas, utilizadas para la "
            "proyección de los factores de mejoramiento"
        ),
    )
    sheet.cell(row=4, column=1, value=sex)
    sheet.cell(row=6, column=1, value="Edad")
    for column, year in enumerate(range(1982, 2017), start=2):
        sheet.cell(row=6, column=column, value=year)
    for row, age in enumerate(range(101), start=7):
        sheet.cell(row=row, column=1, value=age)
        for column in range(2, 37):
            sheet.cell(row=row, column=column, value=0.0)
    sheet.cell(row=110, column=51).fill = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF",
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _member_payloads(
    *,
    rows: Sequence[Mapping[str, object]] | None = None,
    controls: ControlCounts | None = None,
) -> dict[str, bytes]:
    selected_rows = _closed_routing_rows() if rows is None else list(rows)
    selected_controls = {} if controls is None else controls
    payloads = {
        _MICRODATA_MEMBER: _microdata_bytes(selected_rows),
        _INE_MEMBERS[0]: _ine_workbook("HOMBRES"),
        _INE_MEMBERS[1]: _ine_workbook("MUJERES"),
    }
    for table_id, (member, _) in _CONTROL_MEMBERS.items():
        payloads[member] = _mortality_workbook(
            table_id,
            selected_controls.get(table_id, {}),
        )
    return payloads


def _zip_bytes(
    *,
    rows: Sequence[Mapping[str, object]] | None = None,
    controls: ControlCounts | None = None,
    mutation: str | None = None,
) -> bytes:
    payloads = _member_payloads(rows=rows, controls=controls)
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for member in _EXPECTED_MEMBERS:
            if mutation == "missing_member" and member == _INE_MEMBERS[0]:
                continue
            archive.writestr(member, payloads[member])
        if mutation == "extra_member":
            archive.writestr("unexpected/member.txt", b"generated")
        if mutation == "duplicate_member":
            archive.writestr(_MICRODATA_MEMBER, payloads[_MICRODATA_MEMBER])
    return output.getvalue()


def _replace_zip_member(
    archive_raw: bytes,
    member: str,
    payload: bytes,
) -> bytes:
    with ZipFile(BytesIO(archive_raw)) as source:
        payloads = {name: source.read(name) for name in source.namelist()}
    payloads[member] = payload
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as destination:
        for name in sorted(payloads):
            destination.writestr(name, payloads[name])
    return output.getvalue()


def _add_nested_xlsx_member(
    workbook_raw: bytes,
    member: str,
    payload: bytes,
) -> bytes:
    output = BytesIO()
    with (
        ZipFile(BytesIO(workbook_raw)) as source,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            destination.writestr(info, source.read(info.filename))
        destination.writestr(member, payload)
    return output.getvalue()


def _sidecar_document(
    *,
    deltas: Mapping[tuple[str, int], Count] | None = None,
) -> dict[str, Any]:
    selected = {} if deltas is None else deltas
    return {
        "schema_version": "cmf-tm2020-public-original-deviation-v1",
        "descriptor_source_key": _DESCRIPTOR_KEY,
        "descriptor_sha256": _DESCRIPTOR_SHA256,
        "matrix_convention": "public_minus_original",
        "table_order": list(_TABLE_ORDER),
        "age_start": 0,
        "age_end": 110,
        "statistics": ["exposed", "deaths"],
        "rows": [
            {
                "table_id": table_id,
                "age": age,
                "delta_exposed": selected.get((table_id, age), (0, 0))[0],
                "delta_deaths": selected.get((table_id, age), (0, 0))[1],
            }
            for table_id in _TABLE_ORDER
            for age in range(111)
        ],
    }


def _json_bytes(document: Mapping[str, Any]) -> bytes:
    return json.dumps(
        document,
        ensure_ascii=True,
        separators=(",", ":"),
    ).encode()


def _contract(
    deviation_raw: bytes,
    *,
    expected_public_row_count: int = 7,
) -> CmfMortalityConstructionContract:
    return CmfMortalityConstructionContract(
        contract_id=_CONTRACT_ID,
        expected_members=_EXPECTED_MEMBERS,
        descriptor_source_key=_DESCRIPTOR_KEY,
        descriptor_sha256=_DESCRIPTOR_SHA256,
        deviation_source_key=_DEVIATION_KEY,
        deviation_sha256=hashlib.sha256(deviation_raw).hexdigest(),
        table_order=_TABLE_ORDER,
        age_start=0,
        age_end=110,
        statistics=("exposed", "deaths"),
        method_id="sp-tm2020-published-anniversary-v1",
        expected_public_row_count=expected_public_row_count,
    )


def _archive_public_row_count(archive_raw: bytes) -> int:
    with ZipFile(BytesIO(archive_raw)) as archive:
        microdata = archive.read(_MICRODATA_MEMBER)
    return microdata.count(b"\n") - 1


def _parse(
    archive_raw: bytes,
    sidecar_document: Mapping[str, Any],
) -> pd.DataFrame:
    deviation_raw = _json_bytes(sidecar_document)
    return parse_cmf_tm2020_construction(
        BytesIO(archive_raw),
        deviation_raw,
        _contract(
            deviation_raw,
            expected_public_row_count=_archive_public_row_count(archive_raw),
        ),
    ).frame


def _parse_target(
    target: Mapping[str, object],
    counts: Mapping[int, Count],
    *,
    table_id: str,
) -> pd.DataFrame:
    rows = [*_closed_routing_rows(), target]
    archive_raw = _zip_bytes(
        rows=rows,
        controls={table_id: counts},
    )
    return _parse(archive_raw, _sidecar_document())


def _counts(*entries: tuple[int, int, int]) -> dict[int, Count]:
    return {age: (exposed, deaths) for age, exposed, deaths in entries}


def _parser_exception_artifacts(error: BaseException) -> str:
    """Render only parser-owned exception state, never caller-owned fixtures."""

    artifacts: list[str] = []
    pending: list[BaseException] = [error]
    seen: set[int] = set()
    while pending:
        current = pending.pop()
        if id(current) in seen:
            continue
        seen.add(id(current))
        artifacts.extend((str(current), repr(current)))
        traceback = current.__traceback__
        while traceback is not None:
            frame = traceback.tb_frame
            if frame.f_globals.get("__name__") == pensions_module.__name__:
                artifacts.extend(f"{name}={value!r}" for name, value in frame.f_locals.items())
            traceback = traceback.tb_next
        if current.__cause__ is not None:
            pending.append(current.__cause__)
        if current.__context__ is not None:
            pending.append(current.__context__)
    return "\n".join(artifacts)


@pytest.fixture(scope="module")
def zero_release_bytes() -> tuple[bytes, dict[str, Any]]:
    return _zip_bytes(), _sidecar_document()


@pytest.mark.parametrize(
    "mutation",
    ["missing_member", "extra_member", "duplicate_member"],
)
def test_construction_zip_member_multiset_is_exact(
    mutation: str,
) -> None:
    sidecar = _sidecar_document()
    with pytest.raises(DataContractError, match="member"):
        _parse(_zip_bytes(mutation=mutation), sidecar)


def test_changed_child_workbook_header_fails_closed(
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    member = _CONTROL_MEMBERS["CB-H-2020"][0]
    with ZipFile(BytesIO(archive_raw)) as archive:
        workbook_raw = archive.read(member)
    workbook = load_workbook(BytesIO(workbook_raw))
    workbook["qx_brutos"]["B6"] = "changed generated header"
    changed_workbook = BytesIO()
    workbook.save(changed_workbook)
    changed_archive = _replace_zip_member(
        archive_raw,
        member,
        changed_workbook.getvalue(),
    )

    with pytest.raises(DataContractError, match=r"header|qx_brutos"):
        _parse(changed_archive, sidecar)


@pytest.mark.parametrize("cell", ["B3", "B108"])
def test_ine_formatting_only_cells_must_remain_blank(
    cell: str,
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    member = _INE_MEMBERS[0]
    with ZipFile(BytesIO(archive_raw)) as archive:
        workbook_raw = archive.read(member)
    workbook = load_workbook(BytesIO(workbook_raw))
    workbook["TM_históricas"][cell] = 0.123
    changed_workbook = BytesIO()
    workbook.save(changed_workbook)
    changed_archive = _replace_zip_member(
        archive_raw,
        member,
        changed_workbook.getvalue(),
    )

    with pytest.raises(DataContractError, match=r"INE.*blank|formatting"):
        _parse(changed_archive, sidecar)


def test_missing_reviewed_route_combination_fails_closed() -> None:
    rows = _closed_routing_rows()[:-1]
    with pytest.raises(DataContractError):
        _parse(_zip_bytes(rows=rows), _sidecar_document())


def test_invalid_public_row_never_echoes_opaque_token_or_generated_dates() -> None:
    token = "never" + "-echo-this-generated-token"
    generated_date = 20150102
    rows = [
        *_closed_routing_rows(),
        _row(
            9,
            "M",
            "N",
            pension_date=generated_date,
            token=token,
        ),
    ]

    with pytest.raises(DataContractError) as caught:
        _parse(_zip_bytes(rows=rows), _sidecar_document())
    assert token not in str(caught.value)
    assert str(generated_date) not in str(caught.value)


def test_malformed_arrow_error_retains_no_identifier_or_exact_date() -> None:
    token = "never" + "-retain-malformed-id"
    pension_date = 20150102
    birth_date = 19500304
    archive_raw = _zip_bytes()
    with ZipFile(BytesIO(archive_raw)) as archive:
        microdata = archive.read(_MICRODATA_MEMBER)
    malformed = (f"1;{token};2;{pension_date};99991231;F;99;N;{birth_date};0;unexpected\n").encode()
    changed_archive = _replace_zip_member(
        archive_raw,
        _MICRODATA_MEMBER,
        microdata + malformed,
    )
    deviation_raw = _json_bytes(_sidecar_document())

    with pytest.raises(DataContractError) as caught:
        parse_cmf_tm2020_construction(
            BytesIO(changed_archive),
            deviation_raw,
            _contract(deviation_raw),
        )

    artifacts = _parser_exception_artifacts(caught.value)
    assert token not in artifacts
    assert str(pension_date) not in artifacts
    assert str(birth_date) not in artifacts


def test_semantic_error_retains_no_exact_date_in_parser_traceback_locals() -> None:
    marker = 20150102
    rows = [
        *_closed_routing_rows(),
        _row(
            9,
            "M",
            "N",
            pension_date=marker,
            token="not-s" + "elected-by-arrow",
        ),
    ]
    archive_raw = _zip_bytes(rows=rows)
    deviation_raw = _json_bytes(_sidecar_document())

    with pytest.raises(DataContractError) as caught:
        parse_cmf_tm2020_construction(
            BytesIO(archive_raw),
            deviation_raw,
            _contract(deviation_raw),
        )

    assert str(marker) not in _parser_exception_artifacts(caught.value)


def test_nested_xlsx_expansion_budget_fails_before_openpyxl() -> None:
    archive_raw = _zip_bytes()
    member = _CONTROL_MEMBERS["CB-H-2020"][0]
    with ZipFile(BytesIO(archive_raw)) as archive:
        workbook_raw = archive.read(member)
    expanded = _add_nested_xlsx_member(
        workbook_raw,
        "xl/generated-expansion.bin",
        b"x" * 2_000_000,
    )
    changed_archive = _replace_zip_member(archive_raw, member, expanded)

    with pytest.raises(DataContractError, match=r"safety|budget|expansion"):
        _parse(changed_archive, _sidecar_document())


def test_oversized_child_xlsx_fails_before_archive_read(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    member = _CONTROL_MEMBERS["RV-M-2020"][0]
    changed_archive = _replace_zip_member(
        _zip_bytes(),
        member,
        Random(0).randbytes(1_000_001),
    )
    read_members: list[str] = []
    original_read = ZipFile.read

    def recording_read(
        self: ZipFile,
        name: str,
        pwd: bytes | None = None,
    ) -> bytes:
        read_members.append(name)
        return original_read(self, name, pwd)

    monkeypatch.setattr(ZipFile, "read", recording_read)

    with pytest.raises(DataContractError, match=r"safety|budget"):
        _parse(changed_archive, _sidecar_document())
    assert member not in read_members


def test_oversized_deviation_sidecar_fails_before_json_decode() -> None:
    oversized = b" " * 1_000_001

    with pytest.raises(DataContractError, match=r"size|safety|budget"):
        parse_cmf_tm2020_construction(
            BytesIO(_zip_bytes()),
            oversized,
            _contract(oversized),
        )


def test_public_extract_row_count_is_an_exact_contract_field() -> None:
    deviation_raw = _json_bytes(_sidecar_document())
    contract = CmfMortalityConstructionContract(
        contract_id=_CONTRACT_ID,
        expected_members=_EXPECTED_MEMBERS,
        descriptor_source_key=_DESCRIPTOR_KEY,
        descriptor_sha256=_DESCRIPTOR_SHA256,
        deviation_source_key=_DEVIATION_KEY,
        deviation_sha256=hashlib.sha256(deviation_raw).hexdigest(),
        table_order=_TABLE_ORDER,
        age_start=0,
        age_end=110,
        statistics=("exposed", "deaths"),
        method_id="sp-tm2020-published-anniversary-v1",
        expected_public_row_count=2_524_226,
    )

    with pytest.raises(DataContractError, match=r"row count|2,524,226"):
        parse_cmf_tm2020_construction(
            BytesIO(_zip_bytes()),
            deviation_raw,
            contract,
        )


@pytest.mark.parametrize("field", ["delta_exposed", "delta_deaths"])
def test_sidecar_signed_counts_must_fit_int64(field: str) -> None:
    document = _sidecar_document()
    document["rows"][0][field] = 2**63
    deviation_raw = _json_bytes(document)

    with pytest.raises(DataContractError, match=r"integer|int64|range"):
        parse_cmf_tm2020_construction(
            BytesIO(_zip_bytes()),
            deviation_raw,
            _contract(deviation_raw),
        )


def test_sidecar_has_closed_555_row_signed_matrix(
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    frame = _parse(archive_raw, sidecar)

    assert len(sidecar["rows"]) == 555
    assert len(frame) == 555
    assert tuple(frame[["table_id", "age"]].itertuples(index=False, name=None)) == tuple(
        (table_id, age) for table_id in _TABLE_ORDER for age in range(111)
    )


def _mutate_missing_row(document: dict[str, Any]) -> None:
    document["rows"].pop()


def _mutate_extra_row(document: dict[str, Any]) -> None:
    document["rows"].append(
        {
            "table_id": "MI-H-2020",
            "age": 111,
            "delta_exposed": 0,
            "delta_deaths": 0,
        }
    )


def _mutate_duplicate_row(document: dict[str, Any]) -> None:
    document["rows"].append(dict(document["rows"][0]))


def _mutate_table_order(document: dict[str, Any]) -> None:
    document["table_order"] = list(reversed(document["table_order"]))


def _mutate_convention(document: dict[str, Any]) -> None:
    document["matrix_convention"] = "original_minus_public"


def _mutate_noninteger(document: dict[str, Any]) -> None:
    document["rows"][0]["delta_exposed"] = 0.5


def _mutate_unknown_top_level(document: dict[str, Any]) -> None:
    document["unreviewed"] = True


def _mutate_unknown_row_field(document: dict[str, Any]) -> None:
    document["rows"][0]["opaque_fixture_token"] = "must-not-be-accepted"


def _mutate_descriptor_binding(document: dict[str, Any]) -> None:
    document["descriptor_sha256"] = "0" * 64


@pytest.mark.parametrize(
    "mutator",
    [
        _mutate_missing_row,
        _mutate_extra_row,
        _mutate_duplicate_row,
        _mutate_table_order,
        _mutate_convention,
        _mutate_noninteger,
        _mutate_unknown_top_level,
        _mutate_unknown_row_field,
        _mutate_descriptor_binding,
    ],
    ids=lambda mutator: mutator.__name__.removeprefix("_mutate_"),
)
def test_sidecar_shape_and_descriptor_binding_fail_closed(
    zero_release_bytes: tuple[bytes, dict[str, Any]],
    mutator: Callable[[dict[str, Any]], None],
) -> None:
    archive_raw, baseline = zero_release_bytes
    changed = json.loads(json.dumps(baseline))
    mutator(changed)

    with pytest.raises(DataContractError):
        _parse(archive_raw, changed)


def test_sidecar_bytes_must_match_the_reviewed_digest(
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    reviewed_raw = _json_bytes(sidecar)
    changed_raw = reviewed_raw + b"\n"

    with pytest.raises(DataContractError, match=r"hash|SHA|digest"):
        parse_cmf_tm2020_construction(
            BytesIO(archive_raw),
            changed_raw,
            _contract(reviewed_raw),
        )


def test_reconciliation_exposes_only_bounded_aggregate_rows_and_deep_copies(
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    deviation_raw = _json_bytes(sidecar)
    result = parse_cmf_tm2020_construction(
        BytesIO(archive_raw),
        deviation_raw,
        _contract(deviation_raw),
    )
    first = result.frame
    pristine = result.frame

    assert tuple(first.columns) == _AUDIT_COLUMNS
    assert {
        "ID",
        "FEC_VIG",
        "FIN_EXPOSIC",
        "FEC_NAC",
        "FEC_FALL",
    }.isdisjoint(first.columns)
    first.loc[0, "public_exposed"] = 999
    pd.testing.assert_frame_equal(result.frame, pristine)
    assert result.frame is not first


def test_nearest_actuarial_age_controls_both_anniversary_boundaries() -> None:
    frame = _parse_target(
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            birth_date=19500501,
            token="neare" + "st-age-boundary",
        ),
        _counts(
            (65, 1, 0),
            (66, 1, 0),
            (67, 1, 0),
            (68, 1, 0),
        ),
        table_id="CB-H-2020",
    )
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[64, "public_exposed"] == 0
    assert table.loc[65, "public_exposed"] == 1
    assert table.loc[68, "public_exposed"] == 1


def test_study_begins_at_the_2014_anniversary() -> None:
    frame = _parse_target(
        _row(
            3,
            "M",
            "N",
            pension_date=20100101,
            birth_date=19500101,
            token="study" + "-start-boundary",
        ),
        _counts(
            (64, 1, 0),
            (65, 1, 0),
            (66, 1, 0),
            (67, 1, 0),
            (68, 1, 0),
        ),
        table_id="CB-H-2020",
    )
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[63, "public_exposed"] == 0
    assert table.loc[64, "public_exposed"] == 1


@pytest.mark.parametrize(
    ("relationship", "expected_ages"),
    [
        (99, (43, 44)),
        (10, (40, 41, 42, 43, 44)),
    ],
)
def test_three_year_delay_applies_only_to_total_invalid_causants(
    relationship: int,
    expected_ages: tuple[int, ...],
) -> None:
    frame = _parse_target(
        _row(
            4,
            "F",
            "T",
            relationship=relationship,
            pension_date=20140101,
            birth_date=19740101,
            token=f"invalid-delay-{relationship}",
        ),
        {age: (1, 0) for age in expected_ages},
        table_id="MI-M-2020",
    )
    table = frame.loc[frame["table_id"].eq("MI-M-2020")].set_index("age")

    assert tuple(table.index[table["public_exposed"].eq(1)]) == expected_ages


def test_three_year_invalidity_delay_changes_only_the_exposure_start() -> None:
    frame = _parse_target(
        _row(
            4,
            "F",
            "T",
            relationship=99,
            pension_date=20140101,
            birth_date=19740501,
            death_date=20180301,
            token="inval" + "id-delay-does-not-shift-theta",
        ),
        _counts(
            (43, 1, 0),
            (44, 1, 1),
        ),
        table_id="MI-M-2020",
    )
    table = frame.loc[frame["table_id"].eq("MI-M-2020")].set_index("age")

    assert tuple(table.index[table["public_exposed"].eq(1)]) == (43, 44)
    assert tuple(table.index[table["public_deaths"].eq(1)]) == (44,)


@pytest.mark.parametrize(
    ("invalidity", "expected_ages"),
    [
        ("N", (23,)),
        ("P", (23, 24, 25, 26, 27)),
    ],
)
def test_age_24_child_cap_excludes_only_non_invalid_children(
    invalidity: str,
    expected_ages: tuple[int, ...],
) -> None:
    table_number = 1 if invalidity == "N" else 4
    table_id = "B-M-2020" if invalidity == "N" else "MI-M-2020"
    frame = _parse_target(
        _row(
            table_number,
            "F",
            invalidity,
            relationship=30,
            pension_date=20100101,
            birth_date=19910101,
            token=f"child-cap-{invalidity}",
        ),
        {age: (1, 0) for age in expected_ages},
        table_id=table_id,
    )
    table = frame.loc[frame["table_id"].eq(table_id)].set_index("age")

    assert tuple(table.index[table["public_exposed"].eq(1)]) == expected_ages


def test_finite_exposure_end_uses_calendar_year_phi() -> None:
    frame = _parse_target(
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            end_exposure=20171231,
            birth_date=19500101,
            token="calen" + "dar-year-censor",
        ),
        _counts((65, 1, 0), (66, 1, 0)),
        table_id="CB-H-2020",
    )
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[66, "public_exposed"] == 1
    assert table.loc[67, "public_exposed"] == 0


def test_publisher_rows_with_preexisting_pension_or_death_dates_are_not_refiltered() -> None:
    beneficiary = _parse_target(
        _row(
            1,
            "F",
            "N",
            relationship=30,
            pension_date=20100101,
            birth_date=20150101,
            token="benef" + "it-precedes-beneficiary-birth",
        ),
        {age: (1, 0) for age in range(4)},
        table_id="B-M-2020",
    )
    assert tuple(
        beneficiary.loc[
            beneficiary["table_id"].eq("B-M-2020") & beneficiary["public_exposed"].eq(1),
            "age",
        ]
    ) == (0, 1, 2, 3)

    no_interval = _parse(
        _zip_bytes(
            rows=[
                *_closed_routing_rows(),
                _row(
                    3,
                    "M",
                    "N",
                    pension_date=20150101,
                    birth_date=19500101,
                    death_date=20140101,
                    token="death" + "-precedes-pension-start",
                ),
            ]
        ),
        _sidecar_document(),
    )
    assert no_interval["public_exposed"].sum() == 0


def test_fractional_death_age_counts_death_inside_exposed_subgroup() -> None:
    frame = _parse_target(
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            birth_date=19500101,
            death_date=20160501,
            token="fract" + "ional-death-boundary",
        ),
        _counts((65, 1, 0), (66, 1, 1)),
        table_id="CB-H-2020",
    )
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[66, "public_deaths"] == 1
    assert table.loc[67, "public_exposed"] == 0


def test_death_age_uses_rounded_insured_age_plus_time_since_pension() -> None:
    frame = _parse_target(
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            birth_date=19500501,
            death_date=20150301,
            token="publi" + "shed-theta-not-direct-birth-death-age",
        ),
        _counts((65, 1, 1)),
        table_id="CB-H-2020",
    )
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[64, "public_exposed"] == 0
    assert table.loc[64, "public_deaths"] == 0
    assert table.loc[65, "public_exposed"] == 1
    assert table.loc[65, "public_deaths"] == 1


def _signed_fixture() -> tuple[bytes, dict[str, Any]]:
    rows = [
        *_closed_routing_rows(),
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            birth_date=19500101,
            token="signe" + "d-no-death",
        ),
        _row(
            3,
            "M",
            "N",
            pension_date=20150101,
            birth_date=19500101,
            death_date=20160501,
            token="signe" + "d-with-death",
        ),
    ]
    workbook_counts = {
        "CB-H-2020": _counts(
            (65, 1, 0),
            (66, 3, 2),
            (67, 1, 0),
            (68, 1, 0),
        )
    }
    deltas = {
        ("CB-H-2020", 65): (1, 0),
        ("CB-H-2020", 66): (-1, -1),
    }
    return (
        _zip_bytes(rows=rows, controls=workbook_counts),
        _sidecar_document(deltas=deltas),
    )


def test_signed_public_minus_original_counts_reconcile_without_raw_equality() -> None:
    archive_raw, sidecar = _signed_fixture()
    frame = _parse(archive_raw, sidecar)
    table = frame.loc[frame["table_id"].eq("CB-H-2020")].set_index("age")

    assert table.loc[65, "public_exposed"] == 2
    assert table.loc[65, "workbook_exposed"] == 1
    assert table.loc[65, "published_delta_exposed"] == 1
    assert table.loc[66, "public_exposed"] == 2
    assert table.loc[66, "workbook_exposed"] == 3
    assert table.loc[66, "published_delta_exposed"] == -1
    assert table.loc[66, "published_delta_deaths"] == -1
    assert table.loc[66, "public_qx"] == 0.5
    assert table.loc[66, "workbook_qx"] == pytest.approx(2 / 3)
    assert table.loc[66, "public_qx"] != table.loc[66, "workbook_qx"]
    assert frame[["exposed_relation_ok", "deaths_relation_ok", "qx_reconciled"]].all(axis=None)


def test_changed_public_aggregate_count_fails_the_complete_reconciliation() -> None:
    archive_raw, sidecar = _signed_fixture()
    with ZipFile(BytesIO(archive_raw)) as archive:
        microdata = archive.read(_MICRODATA_MEMBER).decode().splitlines()
    changed_microdata = ("\n".join([*microdata, microdata[-1]]) + "\n").encode()
    changed_archive = _replace_zip_member(
        archive_raw,
        _MICRODATA_MEMBER,
        changed_microdata,
    )

    with pytest.raises(
        DataContractError,
        match=r"CB-H-2020.*age.*(?:exposed|deaths)",
    ):
        _parse(changed_archive, sidecar)


@pytest.mark.parametrize("statistic", ["exposed", "deaths"])
def test_reversed_signed_delta_names_table_age_and_statistic(
    statistic: str,
) -> None:
    archive_raw, baseline = _signed_fixture()
    changed = json.loads(json.dumps(baseline))
    row = next(
        item for item in changed["rows"] if item["table_id"] == "CB-H-2020" and item["age"] == 66
    )
    field = f"delta_{statistic}"
    row[field] = -row[field]

    with pytest.raises(
        DataContractError,
        match=rf"CB-H-2020.*66.*{statistic}",
    ):
        _parse(archive_raw, changed)


def _verified_diagnostic_release(
    tmp_path: Path,
    *,
    archive_raw: bytes,
    deviation_raw: bytes,
) -> tuple[CmfMortalityRelease, ReviewedOfficialReleaseRegistry]:
    payloads = {
        _ZIP_KEY: archive_raw,
        _DESCRIPTOR_KEY: _DESCRIPTOR_RAW,
        _DEVIATION_KEY: deviation_raw,
    }
    required_assets: list[RequiredReleaseAsset] = []
    acquired_assets: list[AcquiredAsset] = []
    for source_key, payload in payloads.items():
        path = tmp_path / source_key
        path.write_bytes(payload)
        digest = hashlib.sha256(payload).hexdigest()
        required_assets.append(
            RequiredReleaseAsset(
                source_key=source_key,
                sha256=digest,
                available_at=_NOW,
                released_at=None,
                release_missingness_reason=_MISSING_DATE,
            )
        )
        acquired_assets.append(
            AcquiredAsset(
                key=source_key,
                path=path,
                sha256=digest,
                retrieved_at=_NOW,
            )
        )
    manifest = ReleaseManifest.create(
        schema_version=_RELEASE_SCHEMA,
        assets=tuple(required_assets),
        primary_fact_source_key=_ZIP_KEY,
        identity_registry=IdentityRegistry(),
    )
    verified = verify_release_assets(manifest, acquired_assets)
    release = CmfMortalityRelease(
        assets=verified,
        vintage="TM-2020-construction-generated-test",
        rule_id="NCG_495_SP_NCG_306",
        rule_released_at=date(2023, 2, 24),
        effective_start=date(2023, 7, 1),
        effective_end=None,
        maximum_effective_end=date(2029, 7, 1),
    )
    reviewed_assets = tuple(
        ReviewedOfficialAsset(
            source_key=asset.source_key,
            sha256=asset.sha256,
            released_at=asset.released_at,
            release_missingness_reason=asset.release_missingness_reason,
        )
        for asset in manifest.assets
    )
    entry = ReviewedOfficialReleaseEntry(
        registry_version=_REGISTRY_VERSION,
        registry_key=_REGISTRY_KEY,
        release_schema_version=_RELEASE_SCHEMA,
        primary_fact_source_key=_ZIP_KEY,
        assets=reviewed_assets,
        contract_id=_CONTRACT_ID,
        review_id="generated-test-registry-is-not-production-authorization",
    )
    registry = ReviewedOfficialReleaseRegistry.create(
        registry_version=_REGISTRY_VERSION,
        entries=(entry,),
    )
    return release, registry


@pytest.mark.parametrize("use_injected_registry", [False, True])
def test_unreviewed_production_diagnostic_stays_locked_before_open(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    zero_release_bytes: tuple[bytes, dict[str, Any]],
    use_injected_registry: bool,
) -> None:
    archive_raw, sidecar = zero_release_bytes
    deviation_raw = _json_bytes(sidecar)
    release, test_registry = _verified_diagnostic_release(
        tmp_path,
        archive_raw=archive_raw,
        deviation_raw=deviation_raw,
    )
    opened: list[str] = []

    def forbidden_open(
        self: VerifiedReleaseAssets,
        source_key: str,
    ) -> Any:
        del self
        opened.append(source_key)
        raise AssertionError("authorization must precede opening bytes")

    monkeypatch.setattr(VerifiedReleaseAssets, "open", forbidden_open)
    adapter = (
        CmfMortalityTableAdapter(test_registry)
        if use_injected_registry
        else CmfMortalityTableAdapter()
    )

    with pytest.raises(MissingOfficialDataError):
        adapter.reconcile_anonymized_public_extract(release)
    assert opened == []


def test_injected_registry_and_contract_authorize_exact_diagnostic(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    archive_raw, sidecar = zero_release_bytes
    deviation_raw = _json_bytes(sidecar)
    release, test_registry = _verified_diagnostic_release(
        tmp_path,
        archive_raw=archive_raw,
        deviation_raw=deviation_raw,
    )
    opened: list[str] = []
    original_open = VerifiedReleaseAssets.open

    def recording_open(
        self: VerifiedReleaseAssets,
        source_key: str,
    ) -> Any:
        opened.append(source_key)
        return original_open(self, source_key)

    monkeypatch.setattr(VerifiedReleaseAssets, "open", recording_open)
    adapter = CmfMortalityTableAdapter(
        test_registry,
        construction_contract=_contract(deviation_raw),
    )

    result = adapter.reconcile_anonymized_public_extract(release)

    assert len(result) == 555
    assert set(opened) == {_ZIP_KEY, _DESCRIPTOR_KEY, _DEVIATION_KEY}


def test_release_asset_rebinding_cannot_change_authorized_open_set(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    zero_release_bytes: tuple[bytes, dict[str, Any]],
) -> None:
    approved_raw, sidecar = zero_release_bytes
    deviation_raw = _json_bytes(sidecar)
    approved_dir = tmp_path / "approved"
    unreviewed_dir = tmp_path / "unreviewed"
    approved_dir.mkdir()
    unreviewed_dir.mkdir()
    release, registry = _verified_diagnostic_release(
        approved_dir,
        archive_raw=approved_raw,
        deviation_raw=deviation_raw,
    )
    changed_rows = _closed_routing_rows()
    for index, row in enumerate(changed_rows):
        row["ID"] = f"unreviewed-opaque-{index}"
    unreviewed_raw = _zip_bytes(rows=changed_rows)
    assert hashlib.sha256(unreviewed_raw).digest() != hashlib.sha256(approved_raw).digest()
    unreviewed_release, _ = _verified_diagnostic_release(
        unreviewed_dir,
        archive_raw=unreviewed_raw,
        deviation_raw=deviation_raw,
    )
    approved_assets = release.assets
    unreviewed_assets = unreviewed_release.assets
    opened_from: list[VerifiedReleaseAssets] = []
    original_open = VerifiedReleaseAssets.open

    def rebinding_open(
        self: VerifiedReleaseAssets,
        source_key: str,
    ) -> Any:
        opened_from.append(self)
        if self is approved_assets and source_key == _DESCRIPTOR_KEY:
            object.__setattr__(release, "assets", unreviewed_assets)
        return original_open(self, source_key)

    monkeypatch.setattr(VerifiedReleaseAssets, "open", rebinding_open)
    result = CmfMortalityTableAdapter(
        registry,
        construction_contract=_contract(deviation_raw),
    ).reconcile_anonymized_public_extract(release)

    assert len(result) == 555
    assert opened_from == [approved_assets, approved_assets, approved_assets]
