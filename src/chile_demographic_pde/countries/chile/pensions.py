"""Strict Chilean pension-system mortality data adapters.

TM-2020 base probabilities and improvement factors are regulatory inputs, not
population-mortality observations and not projected probabilities.  The
historical workbook is the canonical fact source.  Construction microdata is
handled only by the aggregate diagnostic implemented below in this module.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import BinaryIO, Final, cast
from zipfile import BadZipFile, ZipFile

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.csv as pacsv
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import HttpUrl

from chile_demographic_pde.core.errors import (
    DataContractError,
    MissingOfficialDataError,
)
from chile_demographic_pde.core.provenance import SourceRef, VariableProvenance
from chile_demographic_pde.countries.chile.official_finance_registry import (
    PRODUCTION_OFFICIAL_RELEASE_REGISTRY_V1,
    ReviewedOfficialReleaseRegistry,
    require_reviewed_official_release,
)
from chile_demographic_pde.data.domain_schemas import (
    REGULATORY_MORTALITY_KEY,
    RegulatoryMortalityBatch,
    validate_regulatory_mortality,
)
from chile_demographic_pde.data.envelope import (
    NormalizedDomainBatch,
    assign_observation_ids,
)
from chile_demographic_pde.data.identity import IdentityRegistry
from chile_demographic_pde.data.releases import VerifiedReleaseAssets
from chile_demographic_pde.data.roles import ObservationDomain

_HISTORICAL_REGISTRY_KEY: Final = "cmf_tm2020_historical"
_HISTORICAL_SOURCE_KEY: Final = "cmf_tm2020_historical_xlsx"
_HISTORICAL_CONTRACT_ID: Final = "cmf-tm2020-historical-workbook-v1"
_RULE_ID: Final = "NCG_495_SP_NCG_306"
_RULE_RELEASED_AT: Final = date(2023, 2, 24)
_EFFECTIVE_START: Final = date(2023, 7, 1)
_MAXIMUM_EFFECTIVE_END: Final = date(2029, 7, 1)
_REFERENCE_YEAR: Final = 2020
_IMPROVEMENT_YEARS: Final = tuple(range(2021, 2037))
_NOTE: Final = "*: A partir del año 2036, los factores de mejoramiento se mantienen constantes"
_CMF_HISTORICAL_URL: Final = HttpUrl(
    "https://www.cmfchile.cl/portal/principal/623/w4-article-99359.html"
)


def _canonical_nonblank(value: object, *, field: str) -> str:
    if (
        type(value) is not str
        or not value
        or value != value.strip()
        or unicodedata.normalize("NFC", value) != value
        or any(unicodedata.category(character).startswith("C") for character in value)
    ):
        raise DataContractError(f"TM-2020 {field} must be canonical nonblank text.")
    return value


def _exact_date(value: object, *, field: str) -> date:
    if type(value) is not date:
        raise DataContractError(f"TM-2020 {field} must be a date without a time.")
    return value


def _finite_number(
    value: object,
    *,
    field: str,
    lower: float,
    upper: float,
) -> float:
    if type(value) not in {int, float}:
        raise DataContractError(f"TM-2020 {field} must be numeric.")
    parsed = float(cast("int | float", value))
    if not math.isfinite(parsed) or not lower <= parsed <= upper:
        raise DataContractError(f"TM-2020 {field} must be finite and in [{lower}, {upper}].")
    return parsed


@dataclass(frozen=True, slots=True)
class CmfMortalityWorkbookBlock:
    """Exact position and meaning of one reviewed 2020 workbook block."""

    sheet_name: str
    dimension: str
    start_column: int
    title: str
    age_start: int
    age_end: int
    note_row: int
    table_id: str
    sex: str
    emit: bool
    duplicate_of: str | None = None

    def __post_init__(self) -> None:
        for field in ("sheet_name", "dimension", "title", "table_id", "sex"):
            _canonical_nonblank(getattr(self, field), field=field)
        if type(self.start_column) is not int or self.start_column <= 0:
            raise DataContractError("TM-2020 block start_column must be positive.")
        if (
            type(self.age_start) is not int
            or type(self.age_end) is not int
            or not 0 <= self.age_start <= self.age_end <= 110
        ):
            raise DataContractError("TM-2020 block age support is invalid.")
        if type(self.note_row) is not int or self.note_row <= 0:
            raise DataContractError("TM-2020 block note_row must be positive.")
        if type(self.emit) is not bool:
            raise DataContractError("TM-2020 block emit marker must be boolean.")
        if self.duplicate_of is not None:
            _canonical_nonblank(self.duplicate_of, field="duplicate_of")
            if self.emit:
                raise DataContractError("TM-2020 duplicate blocks cannot also be emitted.")
        elif not self.emit:
            raise DataContractError("TM-2020 non-emitted blocks must declare their duplicate.")


@dataclass(frozen=True, slots=True)
class CmfMortalityWorkbookContract:
    """Immutable closed contract for the six-sheet historical workbook."""

    contract_id: str
    blocks: tuple[CmfMortalityWorkbookBlock, ...]
    improvement_years: tuple[int, ...] = _IMPROVEMENT_YEARS
    note: str = _NOTE

    def __post_init__(self) -> None:
        _canonical_nonblank(self.contract_id, field="workbook contract_id")
        if type(self.blocks) is not tuple or len(self.blocks) != 6:
            raise DataContractError("TM-2020 historical contract requires exactly six blocks.")
        if any(type(block) is not CmfMortalityWorkbookBlock for block in self.blocks):
            raise DataContractError("TM-2020 historical block type is invalid.")
        sheets = tuple(block.sheet_name for block in self.blocks)
        if len(set(sheets)) != len(sheets):
            raise DataContractError("TM-2020 historical sheet names must be unique.")
        if (
            type(self.improvement_years) is not tuple
            or self.improvement_years != _IMPROVEMENT_YEARS
        ):
            raise DataContractError("TM-2020 improvement years must be exactly 2021 through 2036.")
        if self.note != _NOTE:
            raise DataContractError("TM-2020 historical note is not reviewed.")


CMF_TM2020_HISTORICAL_CONTRACT_V1: Final = CmfMortalityWorkbookContract(
    contract_id=_HISTORICAL_CONTRACT_ID,
    blocks=(
        CmfMortalityWorkbookBlock(
            sheet_name="Vejez-Mujeres",
            dimension="A1:AD95",
            start_column=13,
            title="TABLA RV-2020-MUJERES",
            age_start=20,
            age_end=110,
            note_row=95,
            table_id="RV-M-2020",
            sex="female",
            emit=True,
        ),
        CmfMortalityWorkbookBlock(
            sheet_name="Invalidez-Mujeres",
            dimension="A1:Z115",
            start_column=9,
            title="TABLA MI-2020-MUJERES",
            age_start=0,
            age_end=110,
            note_row=115,
            table_id="MI-M-2020",
            sex="female",
            emit=True,
        ),
        CmfMortalityWorkbookBlock(
            sheet_name="Sobrevivencia-Mujeres",
            dimension="A1:Z115",
            start_column=9,
            title="TABLA B-2020-MUJERES",
            age_start=0,
            age_end=110,
            note_row=115,
            table_id="B-M-2020",
            sex="female",
            emit=True,
        ),
        CmfMortalityWorkbookBlock(
            sheet_name="Vejez-Hombres",
            dimension="A1:AD115",
            start_column=13,
            title="TABLA CB-2020-HOMBRES",
            age_start=0,
            age_end=110,
            note_row=115,
            table_id="CB-H-2020",
            sex="male",
            emit=True,
        ),
        CmfMortalityWorkbookBlock(
            sheet_name="Invalidez-Hombres",
            dimension="A1:Z115",
            start_column=9,
            title="TABLA MI-2020 - HOMBRES",
            age_start=0,
            age_end=110,
            note_row=115,
            table_id="MI-H-2020",
            sex="male",
            emit=True,
        ),
        CmfMortalityWorkbookBlock(
            sheet_name="Sobrevivencia-Hombres",
            dimension="A1:Z115",
            start_column=9,
            title="TABLA CB-2020-HOMBRES",
            age_start=0,
            age_end=110,
            note_row=115,
            table_id="CB-H-2020",
            sex="male",
            emit=False,
            duplicate_of="Vejez-Hombres",
        ),
    ),
)


@dataclass(frozen=True, slots=True)
class ParsedMortalityFact:
    """One source fact before envelope identity and provenance assignment."""

    table_id: str
    age: int
    sex: str
    variable: str
    value: float
    reference_year: int
    improvement_year: int | None


def _rebuild_workbook_contract(
    contract: CmfMortalityWorkbookContract,
) -> CmfMortalityWorkbookContract:
    if type(contract) is not CmfMortalityWorkbookContract:
        raise DataContractError("TM-2020 parser requires a typed workbook contract.")
    try:
        blocks = tuple(
            CmfMortalityWorkbookBlock(
                sheet_name=block.sheet_name,
                dimension=block.dimension,
                start_column=block.start_column,
                title=block.title,
                age_start=block.age_start,
                age_end=block.age_end,
                note_row=block.note_row,
                table_id=block.table_id,
                sex=block.sex,
                emit=block.emit,
                duplicate_of=block.duplicate_of,
            )
            for block in contract.blocks
        )
        return CmfMortalityWorkbookContract(
            contract_id=contract.contract_id,
            blocks=blocks,
            improvement_years=tuple(contract.improvement_years),
            note=contract.note,
        )
    except DataContractError:
        raise
    except Exception:
        raise DataContractError("TM-2020 workbook contract could not be revalidated.") from None


def _parse_historical_block(
    worksheet: object,
    block: CmfMortalityWorkbookBlock,
    contract: CmfMortalityWorkbookContract,
) -> tuple[ParsedMortalityFact, ...]:
    cell = worksheet.cell  # type: ignore[attr-defined]
    dimension = worksheet.calculate_dimension()  # type: ignore[attr-defined]
    if dimension != block.dimension:
        raise DataContractError(f"TM-2020 worksheet {block.sheet_name!r} dimension changed.")
    start = block.start_column
    if cell(1, start).value != block.title:
        raise DataContractError(f"TM-2020 block title changed on {block.sheet_name!r}.")
    headers = tuple(cell(2, start + offset).value for offset in range(3))
    if headers != ("Edad", "qx 2020", "Factores de mejoramiento AAx,t"):
        raise DataContractError(f"TM-2020 block header changed on {block.sheet_name!r}.")
    years = tuple(
        cell(3, start + offset).value for offset in range(2, 2 + len(contract.improvement_years))
    )
    expected_years: tuple[object, ...] = (
        *contract.improvement_years[:-1],
        "2036*",
    )
    if years != expected_years:
        raise DataContractError(f"TM-2020 improvement year header changed on {block.sheet_name!r}.")
    if cell(block.note_row, start).value != contract.note:
        raise DataContractError(f"TM-2020 note changed on {block.sheet_name!r}.")

    facts: list[ParsedMortalityFact] = []
    for row, expected_age in enumerate(
        range(block.age_start, block.age_end + 1),
        start=4,
    ):
        observed_age = cell(row, start).value
        if type(observed_age) is not int or observed_age != expected_age:
            raise DataContractError(f"TM-2020 age support changed on {block.sheet_name!r}.")
        probability = _finite_number(
            cell(row, start + 1).value,
            field=f"qx probability on {block.sheet_name}",
            lower=0.0,
            upper=1.0,
        )
        facts.append(
            ParsedMortalityFact(
                table_id=block.table_id,
                age=expected_age,
                sex=block.sex,
                variable="regulatory_mortality_probability",
                value=probability,
                reference_year=_REFERENCE_YEAR,
                improvement_year=None,
            )
        )
        for offset, year in enumerate(contract.improvement_years, start=2):
            factor = _finite_number(
                cell(row, start + offset).value,
                field=f"improvement factor on {block.sheet_name}",
                lower=0.0,
                upper=1.0,
            )
            facts.append(
                ParsedMortalityFact(
                    table_id=block.table_id,
                    age=expected_age,
                    sex=block.sex,
                    variable="regulatory_mortality_improvement_factor",
                    value=factor,
                    reference_year=_REFERENCE_YEAR,
                    improvement_year=year,
                )
            )
    return tuple(facts)


def _historical_block_cell_values(
    worksheet: object,
    block: CmfMortalityWorkbookBlock,
    contract: CmfMortalityWorkbookContract,
) -> tuple[tuple[object, ...], ...]:
    """Snapshot every cell in one complete reviewed block."""

    cell = worksheet.cell  # type: ignore[attr-defined]
    width = 2 + len(contract.improvement_years)
    return tuple(
        tuple(cell(row, block.start_column + offset).value for offset in range(width))
        for row in range(1, block.note_row + 1)
    )


def parse_cmf_tm2020_historical(
    raw: bytes,
    contract: CmfMortalityWorkbookContract,
) -> tuple[ParsedMortalityFact, ...]:
    """Parse the exact six reviewed blocks without registry or URL access."""

    if type(raw) is not bytes:
        raise DataContractError("TM-2020 historical parser requires immutable bytes.")
    reviewed = _rebuild_workbook_contract(contract)
    try:
        workbook = load_workbook(
            BytesIO(raw),
            read_only=False,
            data_only=True,
        )
    except Exception:
        raise DataContractError(
            "TM-2020 historical workbook is not a readable XLSX document."
        ) from None
    if tuple(workbook.sheetnames) != tuple(block.sheet_name for block in reviewed.blocks):
        raise DataContractError("TM-2020 historical workbook sheet order or names changed.")

    parsed_by_sheet: dict[str, tuple[ParsedMortalityFact, ...]] = {}
    reviewed_cells_by_sheet: dict[str, tuple[tuple[object, ...], ...]] = {}
    for block in reviewed.blocks:
        worksheet = workbook[block.sheet_name]
        parsed_by_sheet[block.sheet_name] = _parse_historical_block(
            worksheet,
            block,
            reviewed,
        )
        reviewed_cells_by_sheet[block.sheet_name] = _historical_block_cell_values(
            worksheet,
            block,
            reviewed,
        )
    for block in reviewed.blocks:
        if block.duplicate_of is not None and (
            parsed_by_sheet[block.sheet_name] != parsed_by_sheet[block.duplicate_of]
            or reviewed_cells_by_sheet[block.sheet_name]
            != reviewed_cells_by_sheet[block.duplicate_of]
        ):
            raise DataContractError(
                "TM-2020 duplicate CB-H block differs from its reviewed source."
            )
    facts = tuple(
        fact
        for block in reviewed.blocks
        if block.emit
        for fact in parsed_by_sheet[block.sheet_name]
    )
    if len(facts) != 9_095:
        raise DataContractError("TM-2020 historical unique fact cardinality is not 9,095.")
    return facts


@dataclass(frozen=True, slots=True)
class CmfMortalityRelease:
    """Verified TM-2020 bytes and their distinct regulatory applicability."""

    assets: VerifiedReleaseAssets
    vintage: str
    rule_id: str
    rule_released_at: date
    effective_start: date
    effective_end: date | None
    maximum_effective_end: date

    def __post_init__(self) -> None:
        if type(self.assets) is not VerifiedReleaseAssets:
            raise DataContractError("TM-2020 release requires verified assets.")
        _canonical_nonblank(self.vintage, field="vintage")
        if self.rule_id != _RULE_ID:
            raise DataContractError("TM-2020 rule identifier changed.")
        if _exact_date(self.rule_released_at, field="rule release date") != (_RULE_RELEASED_AT):
            raise DataContractError("TM-2020 rule release date changed.")
        if _exact_date(self.effective_start, field="effective start") != (_EFFECTIVE_START):
            raise DataContractError("TM-2020 effective start changed.")
        if self.effective_end is not None:
            _exact_date(self.effective_end, field="effective end")
            raise DataContractError("TM-2020 actual effective end must remain open.")
        if (
            _exact_date(
                self.maximum_effective_end,
                field="maximum effective end",
            )
            != _MAXIMUM_EFFECTIVE_END
        ):
            raise DataContractError("TM-2020 maximum effective end changed.")


def _historical_source_refs(release: CmfMortalityRelease) -> tuple[SourceRef, ...]:
    sources: list[SourceRef] = []
    for asset in release.assets.manifest.assets:
        sources.append(
            SourceRef(
                source_key=asset.source_key,
                name="CMF historical mortality tables",
                url=_CMF_HISTORICAL_URL,
                release_date=asset.released_at,
                release_missingness_reason=asset.release_missingness_reason,
                retrieved_at=asset.available_at,
                sha256=asset.sha256,
                vintage=release.vintage,
                provisional=False,
            )
        )
    return tuple(sources)


class CmfMortalityTableAdapter:
    """Authorize and normalize canonical TM-2020 regulatory table facts."""

    def __init__(
        self,
        reviewed_registry: ReviewedOfficialReleaseRegistry = (
            PRODUCTION_OFFICIAL_RELEASE_REGISTRY_V1
        ),
        *,
        construction_contract: CmfMortalityConstructionContract | None = None,
    ) -> None:
        if type(reviewed_registry) is not ReviewedOfficialReleaseRegistry:
            raise MissingOfficialDataError(
                "A typed immutable reviewed official registry is required."
            )
        self._reviewed_registry = reviewed_registry
        self._construction_contract = (
            None
            if construction_contract is None
            else _rebuild_construction_contract(construction_contract)
        )

    def normalize_tables(
        self,
        release: CmfMortalityRelease,
        *,
        identity_registry: IdentityRegistry,
    ) -> RegulatoryMortalityBatch:
        if type(release) is not CmfMortalityRelease:
            raise DataContractError("TM-2020 normalization requires a typed release.")
        if type(identity_registry) is not IdentityRegistry:
            raise DataContractError("TM-2020 normalization requires an IdentityRegistry.")
        authorized = require_reviewed_official_release(
            release,
            _HISTORICAL_REGISTRY_KEY,
            self._reviewed_registry,
        )
        if (
            authorized.contract_id != _HISTORICAL_CONTRACT_ID
            or authorized.primary_fact_source_key != _HISTORICAL_SOURCE_KEY
        ):
            raise MissingOfficialDataError("TM-2020 historical parser contract is not reviewed.")
        with release.assets.open(_HISTORICAL_SOURCE_KEY) as stream:
            raw = stream.read()
        facts = parse_cmf_tm2020_historical(
            raw,
            CMF_TM2020_HISTORICAL_CONTRACT_V1,
        )
        manifest = release.assets.manifest
        primary = manifest.primary_fact_asset
        rows: list[dict[str, object]] = []
        for fact in facts:
            is_probability = fact.variable == "regulatory_mortality_probability"
            period_year = fact.reference_year if is_probability else fact.improvement_year
            if period_year is None:
                raise DataContractError("TM-2020 improvement fact is missing its calendar year.")
            rows.append(
                {
                    "domain": "regulatory_mortality_table",
                    "variable": fact.variable,
                    "value": fact.value,
                    "semantic_kind": "probability" if is_probability else "index",
                    "unit": "dimensionless",
                    "population_basis": ("pensioner" if is_probability else "not_applicable"),
                    "observation_role": "regulatory_input",
                    "parity_scope": "not_applicable",
                    "period_start": date(period_year, 1, 1),
                    "period_end": date(period_year + 1, 1, 1),
                    "source_key": primary.source_key,
                    "vintage": release.vintage,
                    "released_at": primary.released_at,
                    "release_missingness_reason": (
                        None
                        if primary.release_missingness_reason is None
                        else primary.release_missingness_reason.value
                    ),
                    "available_at": manifest.available_at,
                    "provisional": False,
                    "transformation_id": None,
                    "transformation_version": None,
                    "aggregation_rule": (
                        "one published TM-2020 base probability or annual "
                        "improvement factor retained without projection"
                    ),
                    "missingness_reason": "not_applicable_observed",
                    "source_status_code": None,
                    "regulatory_basis": "TM_2020",
                    "mortality_table_id": fact.table_id,
                    "age": fact.age,
                    "sex": fact.sex,
                    "population_segment": "pensioner",
                    "reference_year": fact.reference_year,
                    "improvement_year": fact.improvement_year,
                    "effective_start": release.effective_start,
                    "effective_end": release.effective_end,
                    "maximum_effective_end": release.maximum_effective_end,
                }
            )
        payload = pd.DataFrame(rows)
        payload["improvement_year"] = pd.Series(
            [fact.improvement_year for fact in facts],
            index=payload.index,
            dtype=object,
        )
        caller_snapshot = identity_registry.snapshot()
        staging_registry = IdentityRegistry()
        for record in caller_snapshot.records:
            staging_registry.record(record)
        assigned = assign_observation_ids(
            payload,
            domain=ObservationDomain.REGULATORY_MORTALITY_TABLE,
            domain_key=REGULATORY_MORTALITY_KEY,
            source_identity={
                "publisher": (
                    "Comisión para el Mercado Financiero and Superintendencia de Pensiones"
                ),
                "source_system": "TM-2020 historical mortality workbook",
            },
            manifest=manifest,
            registry=staging_registry,
        )
        validated = validate_regulatory_mortality(assigned)
        provenance = VariableProvenance(
            sources=_historical_source_refs(release),
            release_id=manifest.release_id.value,
            available_at=manifest.available_at,
            observation_start=date(2020, 1, 1),
            observation_end=date(2037, 1, 1),
            dimensions=(
                "mortality_table_id",
                "age",
                "sex",
                "reference_year",
                "improvement_year",
            ),
            aggregation_rules=(
                "published base qx and annual improvement factors retained "
                "as separate regulatory facts without projected qx",
            ),
            missingness_reason="not applicable; reviewed blocks are complete",
            notes=(
                "Reference vintage 2020 is distinct from regulatory "
                "applicability beginning 2023-07-01.",
                "TM-2020 is not population-measure mortality.",
            ),
        )
        batch = NormalizedDomainBatch.create(
            manifest=manifest,
            observations=validated,
            provenance=provenance,
            identities=staging_registry.snapshot(),
        )
        if identity_registry.snapshot() != caller_snapshot:
            raise DataContractError("TM-2020 identity registry changed during normalization.")
        for record in batch.identities.records:
            identity_registry.record(record)
        return batch

    def reconcile_anonymized_public_extract(
        self,
        release: CmfMortalityRelease,
    ) -> AnonymizedPublicExtractReconciliation:
        """Authorize all assets and run the aggregate construction diagnostic."""

        if type(release) is not CmfMortalityRelease:
            raise DataContractError("TM-2020 reconciliation requires a typed release.")
        assets = release.assets
        if type(assets) is not VerifiedReleaseAssets:
            raise MissingOfficialDataError(
                "TM-2020 reconciliation requires verified release assets."
            )
        authorized = require_reviewed_official_release(
            assets.manifest,
            _CONSTRUCTION_REGISTRY_KEY,
            self._reviewed_registry,
        )
        contract = self._construction_contract
        if contract is None:
            raise MissingOfficialDataError(
                "The TM-2020 construction parser contract is not reviewed."
            )
        reviewed = _rebuild_construction_contract(contract)
        manifest = assets.manifest
        assets_by_key = {asset.source_key: asset for asset in manifest.assets}
        if (
            authorized.contract_id != reviewed.contract_id
            or authorized.primary_fact_source_key != _CONSTRUCTION_SOURCE_KEY
            or manifest.primary_fact_source_key != _CONSTRUCTION_SOURCE_KEY
            or set(assets_by_key)
            != {
                _CONSTRUCTION_SOURCE_KEY,
                reviewed.descriptor_source_key,
                reviewed.deviation_source_key,
            }
            or assets_by_key[reviewed.descriptor_source_key].sha256 != reviewed.descriptor_sha256
            or assets_by_key[reviewed.deviation_source_key].sha256 != reviewed.deviation_sha256
        ):
            raise MissingOfficialDataError(
                "The TM-2020 construction parser contract is not reviewed."
            )

        # Entering each verified descriptor rechecks its complete content hash.
        # The PDF is provenance-only; it is deliberately never parsed here.
        with assets.open(reviewed.descriptor_source_key):
            pass
        with assets.open(reviewed.deviation_source_key) as stream:
            deviation_raw = stream.read(_MAX_DEVIATION_SIDECAR_BYTES + 1)
        with assets.open(_CONSTRUCTION_SOURCE_KEY) as stream:
            return parse_cmf_tm2020_construction(
                stream,
                deviation_raw,
                reviewed,
            )


_SHA256: Final = re.compile(r"[0-9a-f]{64}")
_CONSTRUCTION_REGISTRY_KEY: Final = "cmf_tm2020_construction_diagnostic"
_CONSTRUCTION_SOURCE_KEY: Final = "cmf_tm2020_construction_zip"
_MAX_DEVIATION_SIDECAR_BYTES: Final = 1_000_000
_MAX_CONSTRUCTION_XLSX_BYTES: Final = 1_000_000
_MAX_CONSTRUCTION_XLSX_PARTS: Final = 100
_MAX_CONSTRUCTION_XLSX_PART_BYTES: Final = 500_000
_MAX_CONSTRUCTION_XLSX_EXPANDED_BYTES: Final = 1_000_000
_MAX_CONSTRUCTION_XLSX_EXPANSION_RATIO: Final = 50.0
_INT64_MIN: Final = -(2**63)
_INT64_MAX: Final = 2**63 - 1
_CONSTRUCTION_MEMBERS: Final = tuple(
    sorted(
        (
            "bd/bd_publica.txt",
            "tablas_ine/tablas_ine-hombres.xlsx",
            "tablas_ine/tablas_ine-mujeres.xlsx",
            "tablas_mortalidad/b 2020-mujeres.xlsx",
            "tablas_mortalidad/cb 2020-hombres.xlsx",
            "tablas_mortalidad/mi 2020-hombres.xlsx",
            "tablas_mortalidad/mi 2020-mujeres.xlsx",
            "tablas_mortalidad/rv 2020-mujeres.xlsx",
        )
    )
)
_CONSTRUCTION_TABLE_ORDER: Final = (
    "RV-M-2020",
    "CB-H-2020",
    "B-M-2020",
    "MI-M-2020",
    "MI-H-2020",
)
_MICRODATA_COLUMNS: Final = (
    "ORIGEN",
    "ID",
    "TABLA",
    "FEC_VIG",
    "FIN_EXPOSIC",
    "SEXO",
    "COD_REL",
    "BEN_COD_INV",
    "FEC_NAC",
    "FEC_FALL",
)
_ANALYTICAL_COLUMNS: Final = tuple(column for column in _MICRODATA_COLUMNS if column != "ID")
_ROUTE_COMBINATIONS: Final = frozenset(
    {
        (1, "F", "N"),
        (2, "F", "N"),
        (3, "M", "N"),
        (4, "F", "P"),
        (4, "F", "T"),
        (4, "M", "P"),
        (4, "M", "T"),
    }
)
_RELATIONSHIPS: Final = frozenset({10, 11, 20, 21, 30, 35, 41, 42, 50, 51, 52, 99})
_CONTROL_MEMBERS: Final = (
    (
        "tablas_mortalidad/rv 2020-mujeres.xlsx",
        "RV-M-2020",
        "RV-2020, Mujeres",
    ),
    (
        "tablas_mortalidad/cb 2020-hombres.xlsx",
        "CB-H-2020",
        "CB-2020, Hombres",
    ),
    (
        "tablas_mortalidad/b 2020-mujeres.xlsx",
        "B-M-2020",
        "B-2020, Mujeres",
    ),
    (
        "tablas_mortalidad/mi 2020-mujeres.xlsx",
        "MI-M-2020",
        "MI-2020, Mujeres",
    ),
    (
        "tablas_mortalidad/mi 2020-hombres.xlsx",
        "MI-H-2020",
        "MI-2020, Hombres",
    ),
)
_AUDIT_COLUMNS: Final = (
    "table_id",
    "age",
    "public_exposed",
    "public_deaths",
    "public_qx",
    "workbook_exposed",
    "workbook_deaths",
    "workbook_qx",
    "published_delta_exposed",
    "published_delta_deaths",
    "exposed_relation_ok",
    "deaths_relation_ok",
    "qx_reconciled",
)


@dataclass(frozen=True, slots=True)
class CmfMortalityConstructionContract:
    """Closed aggregate-only parser contract for one descriptor-bound release."""

    contract_id: str
    expected_members: tuple[str, ...]
    descriptor_source_key: str
    descriptor_sha256: str
    deviation_source_key: str
    deviation_sha256: str
    table_order: tuple[str, ...]
    age_start: int
    age_end: int
    statistics: tuple[str, ...]
    method_id: str
    expected_public_row_count: int

    def __post_init__(self) -> None:
        for field in (
            "contract_id",
            "descriptor_source_key",
            "deviation_source_key",
            "method_id",
        ):
            _canonical_nonblank(getattr(self, field), field=field)
        for field in ("descriptor_sha256", "deviation_sha256"):
            value = getattr(self, field)
            if type(value) is not str or _SHA256.fullmatch(value) is None:
                raise DataContractError(
                    f"TM-2020 construction {field} must be a lowercase SHA-256."
                )
        if (
            type(self.expected_members) is not tuple
            or self.expected_members != _CONSTRUCTION_MEMBERS
        ):
            raise DataContractError("TM-2020 construction member contract is not exact.")
        if type(self.table_order) is not tuple or self.table_order != _CONSTRUCTION_TABLE_ORDER:
            raise DataContractError("TM-2020 construction table order is not reviewed.")
        if (
            type(self.age_start) is not int
            or type(self.age_end) is not int
            or (self.age_start, self.age_end) != (0, 110)
        ):
            raise DataContractError("TM-2020 construction age support must be 0 through 110.")
        if type(self.statistics) is not tuple or self.statistics != (
            "exposed",
            "deaths",
        ):
            raise DataContractError("TM-2020 construction statistics must be exposed and deaths.")
        if self.method_id != "sp-tm2020-published-anniversary-v1":
            raise DataContractError(
                "TM-2020 construction method is not the reviewed anniversary method."
            )
        if (
            type(self.expected_public_row_count) is not int
            or not 1 <= self.expected_public_row_count <= 10_000_000
        ):
            raise DataContractError("TM-2020 public extract expected row count is invalid.")


@dataclass(frozen=True, slots=True, init=False, eq=False, repr=False)
class AnonymizedPublicExtractReconciliation:
    """Deep-copy-safe aggregate reconciliation with no individual fields."""

    _frame: pd.DataFrame

    def __init__(self, *_: object, **__: object) -> None:
        raise TypeError("AnonymizedPublicExtractReconciliation is parser-created.")

    @classmethod
    def _create(
        cls,
        frame: pd.DataFrame,
    ) -> AnonymizedPublicExtractReconciliation:
        if not isinstance(frame, pd.DataFrame):
            raise DataContractError("TM-2020 reconciliation requires an aggregate DataFrame.")
        if tuple(frame.columns) != _AUDIT_COLUMNS or len(frame) != 555:
            raise DataContractError("TM-2020 reconciliation aggregate shape is invalid.")
        result = object.__new__(cls)
        object.__setattr__(result, "_frame", frame.copy(deep=True))
        return result

    @property
    def frame(self) -> pd.DataFrame:
        """Return a detached aggregate-only view."""

        return self._frame.copy(deep=True)

    def __len__(self) -> int:
        return len(self._frame)


def _rebuild_construction_contract(
    contract: CmfMortalityConstructionContract,
) -> CmfMortalityConstructionContract:
    if type(contract) is not CmfMortalityConstructionContract:
        raise DataContractError("TM-2020 construction parser requires a typed contract.")
    try:
        return CmfMortalityConstructionContract(
            contract_id=contract.contract_id,
            expected_members=tuple(contract.expected_members),
            descriptor_source_key=contract.descriptor_source_key,
            descriptor_sha256=contract.descriptor_sha256,
            deviation_source_key=contract.deviation_source_key,
            deviation_sha256=contract.deviation_sha256,
            table_order=tuple(contract.table_order),
            age_start=contract.age_start,
            age_end=contract.age_end,
            statistics=tuple(contract.statistics),
            method_id=contract.method_id,
            expected_public_row_count=contract.expected_public_row_count,
        )
    except DataContractError:
        raise
    except Exception:
        raise DataContractError("TM-2020 construction contract could not be revalidated.") from None


def _duplicate_rejecting_object(
    pairs: list[tuple[str, object]],
) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError("duplicate JSON key")
        result[key] = value
    return result


def _parse_deviation_sidecar(
    raw: bytes,
    contract: CmfMortalityConstructionContract,
) -> np.ndarray:
    if type(raw) is not bytes:
        raise DataContractError("TM-2020 deviation sidecar requires immutable bytes.")
    if len(raw) > _MAX_DEVIATION_SIDECAR_BYTES:
        raise DataContractError("TM-2020 deviation sidecar exceeds its safety size budget.")
    if hashlib.sha256(raw).hexdigest() != contract.deviation_sha256:
        raise DataContractError("TM-2020 deviation sidecar SHA-256 digest changed.")
    try:
        document = json.loads(
            raw.decode("utf-8", errors="strict"),
            object_pairs_hook=_duplicate_rejecting_object,
            parse_constant=lambda _value: (_ for _ in ()).throw(ValueError("non-finite JSON")),
        )
    except (UnicodeDecodeError, json.JSONDecodeError, TypeError, ValueError):
        raise DataContractError("TM-2020 deviation sidecar is not strict UTF-8 JSON.") from None
    root_fields = {
        "schema_version",
        "descriptor_source_key",
        "descriptor_sha256",
        "matrix_convention",
        "table_order",
        "age_start",
        "age_end",
        "statistics",
        "rows",
    }
    if type(document) is not dict or set(document) != root_fields:
        raise DataContractError("TM-2020 deviation sidecar root fields are not closed.")
    root = cast(dict[str, object], document)
    if (
        root["schema_version"] != "cmf-tm2020-public-original-deviation-v1"
        or root["descriptor_source_key"] != contract.descriptor_source_key
        or root["descriptor_sha256"] != contract.descriptor_sha256
        or root["matrix_convention"] != "public_minus_original"
        or root["table_order"] != list(contract.table_order)
        or type(root["age_start"]) is not int
        or root["age_start"] != contract.age_start
        or type(root["age_end"]) is not int
        or root["age_end"] != contract.age_end
        or root["statistics"] != list(contract.statistics)
    ):
        raise DataContractError("TM-2020 deviation sidecar descriptor binding changed.")
    rows = root["rows"]
    if type(rows) is not list or len(rows) != 555:
        raise DataContractError("TM-2020 deviation sidecar must contain exactly 555 rows.")
    expected_pairs = tuple(
        (table_id, age)
        for table_id in contract.table_order
        for age in range(contract.age_start, contract.age_end + 1)
    )
    deviations = np.zeros((5, 111, 2), dtype=np.int64)
    observed_pairs: list[tuple[str, int]] = []
    row_fields = {
        "table_id",
        "age",
        "delta_exposed",
        "delta_deaths",
    }
    table_index = {table_id: index for index, table_id in enumerate(contract.table_order)}
    for item in rows:
        if type(item) is not dict or set(item) != row_fields:
            raise DataContractError("TM-2020 deviation sidecar row fields are not closed.")
        row = cast(dict[str, object], item)
        table_id = row["table_id"]
        age = row["age"]
        exposed = row["delta_exposed"]
        deaths = row["delta_deaths"]
        if (
            type(table_id) is not str
            or table_id not in table_index
            or type(age) is not int
            or not contract.age_start <= age <= contract.age_end
            or type(exposed) is not int
            or type(deaths) is not int
            or not _INT64_MIN <= exposed <= _INT64_MAX
            or not _INT64_MIN <= deaths <= _INT64_MAX
        ):
            raise DataContractError("TM-2020 deviation sidecar contains a noncanonical int64 row.")
        observed_pairs.append((table_id, age))
        deviations[table_index[table_id], age] = (exposed, deaths)
    if tuple(observed_pairs) != expected_pairs:
        raise DataContractError(
            "TM-2020 deviation sidecar rows are missing, duplicated, or reordered."
        )
    return deviations


def _strict_integral(value: object, *, field: str) -> int:
    if type(value) is not int or not 0 <= value <= _INT64_MAX:
        raise DataContractError(
            f"TM-2020 construction {field} must be a nonnegative int64 integer."
        )
    return value


def _validate_construction_xlsx_budget(raw: bytes) -> None:
    if type(raw) is not bytes or len(raw) > _MAX_CONSTRUCTION_XLSX_BYTES:
        raise DataContractError("TM-2020 construction nested XLSX safety budget changed.")
    try:
        with ZipFile(BytesIO(raw)) as workbook:
            infos = workbook.infolist()
    except (BadZipFile, OSError, TypeError, ValueError):
        raise DataContractError("TM-2020 construction child workbook is not readable.") from None
    names = tuple(info.filename for info in infos)
    if (
        len(infos) > _MAX_CONSTRUCTION_XLSX_PARTS
        or len(names) != len(set(names))
        or sum(info.file_size for info in infos) > _MAX_CONSTRUCTION_XLSX_EXPANDED_BYTES
        or any(
            info.flag_bits & 0x1
            or info.file_size < 0
            or info.file_size > _MAX_CONSTRUCTION_XLSX_PART_BYTES
            or info.compress_size < 0
            or (info.file_size > 0 and info.compress_size == 0)
            or (
                info.compress_size > 0
                and info.file_size / info.compress_size > _MAX_CONSTRUCTION_XLSX_EXPANSION_RATIO
            )
            for info in infos
        )
    ):
        raise DataContractError("TM-2020 construction nested XLSX safety budget changed.")


def _parse_control_workbook(
    raw: bytes,
    *,
    member: str,
    table_id: str,
    reviewed_sheet: str,
) -> tuple[np.ndarray, np.ndarray]:
    _validate_construction_xlsx_budget(raw)
    try:
        stored_workbook = load_workbook(
            BytesIO(raw),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        stored_sheets = tuple(stored_workbook.sheetnames)
        stored_dimensions = {
            sheet_name: stored_workbook[sheet_name].calculate_dimension()
            for sheet_name in stored_sheets
        }
        stored_workbook.close()
        workbook = load_workbook(
            BytesIO(raw),
            read_only=False,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        raise DataContractError("TM-2020 construction child workbook is not readable.") from None
    if (
        stored_sheets != ("qx_brutos", reviewed_sheet)
        or tuple(workbook.sheetnames) != stored_sheets
    ):
        raise DataContractError(f"TM-2020 {table_id} child workbook sheet contract changed.")
    controls = workbook["qx_brutos"]
    if stored_dimensions["qx_brutos"] != "B2:K117":
        raise DataContractError(f"TM-2020 {table_id} qx_brutos dimension changed.")
    if tuple(controls.cell(6, column).value for column in range(2, 6)) != (
        "Edad",
        "Expuestos",
        "Muertos",
        "qx_bruto",
    ):
        raise DataContractError(f"TM-2020 {table_id} qx_brutos header changed.")
    counts = np.zeros((111, 2), dtype=np.int64)
    ratios = np.full(111, np.nan, dtype=np.float64)
    for age in range(111):
        row = age + 7
        if type(controls.cell(row, 2).value) is not int or (controls.cell(row, 2).value != age):
            raise DataContractError(f"TM-2020 {table_id} qx_brutos age support changed.")
        exposed = _strict_integral(
            controls.cell(row, 3).value,
            field=f"{table_id} exposed",
        )
        deaths = _strict_integral(
            controls.cell(row, 4).value,
            field=f"{table_id} deaths",
        )
        if deaths > exposed:
            raise DataContractError(f"TM-2020 {table_id} deaths exceed exposed at age {age}.")
        ratio_value = controls.cell(row, 5).value
        if exposed == 0:
            if ratio_value is not None:
                raise DataContractError(f"TM-2020 {table_id} zero exposure has a qx value.")
        else:
            ratio = _finite_number(
                ratio_value,
                field=f"{table_id} workbook qx",
                lower=0.0,
                upper=1.0,
            )
            if not math.isclose(
                ratio,
                deaths / exposed,
                rel_tol=1e-12,
                abs_tol=1e-15,
            ):
                raise DataContractError(
                    f"TM-2020 {table_id} workbook qx ratio changed at age {age}."
                )
            ratios[age] = ratio
        if any(controls.cell(row, column).value is not None for column in range(6, 12)):
            raise DataContractError(f"TM-2020 {table_id} qx_brutos blank columns changed.")
        counts[age] = (exposed, deaths)

    reviewed = workbook[reviewed_sheet]
    age_start = 20 if table_id == "RV-M-2020" else 0
    note_row = 104 if table_id == "RV-M-2020" else 124
    expected_dimension = "A2:W104" if table_id == "RV-M-2020" else "A2:W124"
    if stored_dimensions[reviewed_sheet] != expected_dimension:
        raise DataContractError(f"TM-2020 {table_id} reviewed table dimension changed.")
    if tuple(reviewed.cell(7, column).value for column in range(1, 4)) != (
        "Edad",
        "Tasas de mortalidad qx",
        "Factores de mejoramiento AAx,t",
    ):
        raise DataContractError(f"TM-2020 {table_id} reviewed table header changed.")
    expected_years: tuple[object, ...] = (*range(2016, 2036), "2036*")
    if tuple(reviewed.cell(8, column).value for column in range(3, 24)) != (expected_years):
        raise DataContractError(f"TM-2020 {table_id} reviewed improvement years changed.")
    for row, age in enumerate(range(age_start, 111), start=9):
        if type(reviewed.cell(row, 1).value) is not int or (reviewed.cell(row, 1).value != age):
            raise DataContractError(f"TM-2020 {table_id} reviewed age support changed.")
        _finite_number(
            reviewed.cell(row, 2).value,
            field=f"{table_id} reviewed qx",
            lower=0.0,
            upper=1.0,
        )
        for column in range(3, 24):
            _finite_number(
                reviewed.cell(row, column).value,
                field=f"{table_id} reviewed improvement factor",
                lower=0.0,
                upper=1.0,
            )
    if reviewed.cell(note_row, 1).value != _NOTE:
        raise DataContractError(f"TM-2020 {table_id} reviewed note changed.")
    del member
    return counts, ratios


def _validate_ine_workbook(raw: bytes, *, expected_sex: str) -> None:
    _validate_construction_xlsx_budget(raw)
    try:
        stored_workbook = load_workbook(
            BytesIO(raw),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        stored_sheets = tuple(stored_workbook.sheetnames)
        stored_dimension = stored_workbook["TM_históricas"].calculate_dimension()
        stored_workbook.close()
        workbook = load_workbook(
            BytesIO(raw),
            read_only=False,
            data_only=True,
            keep_links=False,
        )
    except Exception:
        raise DataContractError("TM-2020 historical INE workbook is not readable.") from None
    if stored_sheets != ("TM_históricas",) or tuple(workbook.sheetnames) != stored_sheets:
        raise DataContractError("TM-2020 historical INE sheet contract changed.")
    sheet = workbook["TM_históricas"]
    if stored_dimension != "A2:AY110":
        raise DataContractError("TM-2020 historical INE worksheet dimension changed.")
    if (
        sheet["A2"].value
        != (
            "Tablas de mortalidad históricas, utilizadas para la "
            "proyección de los factores de mejoramiento"
        )
        or sheet["A4"].value != expected_sex
        or sheet["A6"].value != "Edad"
    ):
        raise DataContractError("TM-2020 historical INE header changed.")
    if tuple(sheet.cell(6, column).value for column in range(2, 37)) != tuple(range(1982, 2017)):
        raise DataContractError("TM-2020 historical INE years must be exactly 1982 through 2016.")
    for row, age in enumerate(range(101), start=7):
        if type(sheet.cell(row, 1).value) is not int or (sheet.cell(row, 1).value != age):
            raise DataContractError("TM-2020 historical INE age support must be 0 through 100.")
        for column in range(2, 37):
            _finite_number(
                sheet.cell(row, column).value,
                field="historical INE probability",
                lower=0.0,
                upper=1.0,
            )
    if any(
        sheet.cell(row, column).value is not None
        for row in range(2, 111)
        for column in range(1, 52)
        if (
            not (row in (2, 4) and column == 1)
            and not (row == 6 and 1 <= column <= 36)
            and not (7 <= row <= 107 and 1 <= column <= 36)
        )
    ):
        raise DataContractError("TM-2020 historical INE formatting-only cells became populated.")


def _valid_yyyymmdd(values: np.ndarray, *, allow_zero: bool, allow_open: bool) -> bool:
    numbers = values.astype(np.int64, copy=False)
    sentinel = np.zeros(len(numbers), dtype=bool)
    if allow_zero:
        sentinel |= numbers == 0
    if allow_open:
        sentinel |= numbers == 99_991_231
    candidates = numbers[~sentinel]
    if candidates.size == 0:
        return True
    years = candidates // 10_000
    months = (candidates // 100) % 100
    days = candidates % 100
    leap = (years % 4 == 0) & ((years % 100 != 0) | (years % 400 == 0))
    month_days = np.asarray(
        (0, 31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31),
        dtype=np.int64,
    )
    valid_month = (months >= 1) & (months <= 12)
    safe_month = np.clip(months, 1, 12)
    maximum = month_days[safe_month] + ((safe_month == 2) & leap)
    return bool(
        np.all((years >= 1) & (years <= 9999) & valid_month & (days >= 1) & (days <= maximum))
    )


def _route_table(table: np.ndarray, sex: np.ndarray) -> np.ndarray:
    return np.select(
        [
            table == 2,
            table == 3,
            table == 1,
            (table == 4) & (sex == "F"),
            (table == 4) & (sex == "M"),
        ],
        [0, 1, 2, 3, 4],
        default=-1,
    ).astype(np.int64)


def _aggregate_public_extract(
    archive: ZipFile,
    *,
    expected_row_count: int,
) -> np.ndarray:
    differences = np.zeros((5, 112), dtype=np.int64)
    deaths_by_age = np.zeros((5, 111), dtype=np.int64)
    observed_combinations: set[tuple[int, str, str]] = set()
    observed_row_count = 0
    with archive.open("bd/bd_publica.txt") as header_stream:
        try:
            header = (
                header_stream.readline(4096)
                .decode(
                    "utf-8",
                    errors="strict",
                )
                .rstrip("\r\n")
            )
        except UnicodeDecodeError:
            raise DataContractError("TM-2020 public extract header is not strict UTF-8.") from None
    if header != ";".join(_MICRODATA_COLUMNS):
        raise DataContractError("TM-2020 public extract header changed.")

    column_types = {
        "ORIGEN": pa.int32(),
        "TABLA": pa.int32(),
        "FEC_VIG": pa.int32(),
        "FIN_EXPOSIC": pa.int32(),
        "SEXO": pa.string(),
        "COD_REL": pa.int32(),
        "BEN_COD_INV": pa.string(),
        "FEC_NAC": pa.int32(),
        "FEC_FALL": pa.int32(),
    }
    try:
        with archive.open("bd/bd_publica.txt") as microdata:
            reader = pacsv.open_csv(  # type: ignore[attr-defined]
                microdata,
                read_options=pacsv.ReadOptions(  # type: ignore[attr-defined]
                    block_size=1 << 24
                ),
                parse_options=pacsv.ParseOptions(  # type: ignore[attr-defined]
                    delimiter=";"
                ),
                convert_options=pacsv.ConvertOptions(  # type: ignore[attr-defined]
                    column_types=column_types,
                    include_columns=list(_ANALYTICAL_COLUMNS),
                ),
            )
            for batch in reader:
                observed_row_count += batch.num_rows
                if observed_row_count > expected_row_count:
                    raise DataContractError(
                        "TM-2020 public extract row count exceeds its "
                        f"reviewed value of {expected_row_count:,}."
                    )
                if batch.num_rows == 0:
                    continue
                if any(batch.column(name).null_count for name in _ANALYTICAL_COLUMNS):
                    raise DataContractError(
                        "TM-2020 public extract contains missing analytical fields."
                    )
                values = {name: np.asarray(batch.column(name)) for name in _ANALYTICAL_COLUMNS}
                origin = values["ORIGEN"].astype(np.int64)
                table = values["TABLA"].astype(np.int64)
                sex = values["SEXO"].astype(str)
                relationship = values["COD_REL"].astype(np.int64)
                invalidity = values["BEN_COD_INV"].astype(str)
                if (
                    not np.isin(origin, (1, 2)).all()
                    or not np.isin(
                        relationship,
                        tuple(_RELATIONSHIPS),
                    ).all()
                ):
                    raise DataContractError("TM-2020 public extract enum domain changed.")
                combinations = set(
                    zip(
                        table.tolist(),
                        sex.tolist(),
                        invalidity.tolist(),
                        strict=True,
                    )
                )
                if not combinations <= _ROUTE_COMBINATIONS:
                    raise DataContractError("TM-2020 public extract route combination changed.")
                observed_combinations.update(combinations)

                pension = values["FEC_VIG"].astype(np.int64)
                exposure_end = values["FIN_EXPOSIC"].astype(np.int64)
                birth = values["FEC_NAC"].astype(np.int64)
                death = values["FEC_FALL"].astype(np.int64)
                if (
                    not _valid_yyyymmdd(pension, allow_zero=False, allow_open=False)
                    or not _valid_yyyymmdd(
                        exposure_end,
                        allow_zero=False,
                        allow_open=True,
                    )
                    or not _valid_yyyymmdd(birth, allow_zero=False, allow_open=False)
                    or not _valid_yyyymmdd(death, allow_zero=True, allow_open=False)
                ):
                    raise DataContractError(
                        "TM-2020 public extract contains an invalid date encoding."
                    )
                has_death = death != 0
                has_end = exposure_end != 99_991_231

                pension_year = pension // 10_000
                pension_month = (pension // 100) % 100
                pension_day = pension % 100
                birth_year = birth // 10_000
                birth_month = (birth // 100) % 100
                birth_day = birth % 100
                exact_age = (
                    pension_year
                    - birth_year
                    + (pension_month - birth_month) / 12.0
                    + (pension_day - birth_day) / 365.25
                )
                insured_age = np.floor(exact_age + 0.5).astype(np.int64)
                virtual_birth_year = pension_year - insured_age
                delayed = (invalidity == "T") & (relationship == 99)
                lower_age = (
                    np.maximum(
                        2014,
                        np.where(delayed, pension_year + 3, pension_year),
                    )
                    - virtual_birth_year
                )
                upper_age = 2019 - virtual_birth_year - 1

                theta = np.zeros(len(table), dtype=np.float64)
                death_year = death // 10_000
                death_month = (death // 100) % 100
                death_day = death % 100
                theta[has_death] = (
                    insured_age[has_death]
                    + death_year[has_death]
                    - pension_year[has_death]
                    + (death_month[has_death] - pension_month[has_death]) / 12.0
                    + (death_day[has_death] - pension_day[has_death]) / 365.25
                )
                upper_age = np.minimum(
                    upper_age,
                    np.where(
                        has_death,
                        np.ceil(theta).astype(np.int64) - 1,
                        10_000,
                    ),
                )
                censor_age = exposure_end // 10_000 - virtual_birth_year - 1
                upper_age = np.minimum(
                    upper_age,
                    np.where(has_end, censor_age, 10_000),
                )
                capped_child = np.isin(relationship, (30, 35)) & (invalidity == "N")
                upper_age = np.minimum(
                    upper_age,
                    np.where(capped_child, 23, 10_000),
                )

                table_index = _route_table(table, sex)
                lower_age = np.maximum(lower_age, 0)
                upper_age = np.minimum(upper_age, 110)
                valid = (table_index >= 0) & (lower_age <= upper_age)
                flat_lower = table_index[valid] * 112 + lower_age[valid]
                flat_after = table_index[valid] * 112 + upper_age[valid] + 1
                np.add.at(differences.ravel(), flat_lower, 1)
                np.add.at(differences.ravel(), flat_after, -1)

                death_age = np.ceil(theta).astype(np.int64) - 1
                valid_death = (
                    valid
                    & has_death
                    & (death_age >= lower_age)
                    & (death_age <= upper_age)
                    & (death_age >= 0)
                    & (death_age <= 110)
                )
                flat_death = table_index[valid_death] * 111 + death_age[valid_death]
                np.add.at(deaths_by_age.ravel(), flat_death, 1)
    except DataContractError:
        raise
    except Exception:
        raise DataContractError(
            "TM-2020 public extract violates its closed streaming contract."
        ) from None
    if observed_combinations != _ROUTE_COMBINATIONS:
        raise DataContractError("TM-2020 public extract is missing a reviewed route combination.")
    if observed_row_count != expected_row_count:
        raise DataContractError(
            "TM-2020 public extract row count differs from its "
            f"reviewed value of {expected_row_count:,}."
        )
    return np.stack(
        (
            np.cumsum(differences[:, :111], axis=1),
            deaths_by_age,
        ),
        axis=-1,
    )


def _parse_cmf_tm2020_construction(
    raw: BinaryIO,
    deviation_raw: bytes,
    contract: CmfMortalityConstructionContract,
) -> AnonymizedPublicExtractReconciliation:
    reviewed = _rebuild_construction_contract(contract)
    deviations = _parse_deviation_sidecar(deviation_raw, reviewed)
    if not hasattr(raw, "read") or not hasattr(raw, "seek"):
        raise DataContractError("TM-2020 construction parser requires a seekable binary stream.")
    try:
        archive = ZipFile(raw)
    except (BadZipFile, OSError, TypeError, ValueError):
        raise DataContractError("TM-2020 construction release is not a readable ZIP.") from None
    with archive:
        infos = archive.infolist()
        names = tuple(sorted(info.filename for info in infos))
        if len(infos) != len(reviewed.expected_members) or names != (reviewed.expected_members):
            raise DataContractError("TM-2020 construction ZIP member multiset changed.")
        if (
            any(
                info.flag_bits & 0x1
                or info.file_size < 0
                or info.file_size > 200_000_000
                or info.compress_size < 0
                or (info.compress_size > 0 and info.file_size / info.compress_size > 100.0)
                for info in infos
            )
            or sum(info.file_size for info in infos) > 300_000_000
        ):
            raise DataContractError("TM-2020 construction ZIP member safety contract changed.")
        infos_by_name = {info.filename: info for info in infos}
        workbook_members = (
            *(member for member, _, _ in _CONTROL_MEMBERS),
            "tablas_ine/tablas_ine-hombres.xlsx",
            "tablas_ine/tablas_ine-mujeres.xlsx",
        )
        if any(
            infos_by_name[member].file_size > _MAX_CONSTRUCTION_XLSX_BYTES
            or infos_by_name[member].compress_size > _MAX_CONSTRUCTION_XLSX_BYTES
            for member in workbook_members
        ):
            raise DataContractError("TM-2020 construction nested XLSX safety budget changed.")

        controls = np.zeros((5, 111, 2), dtype=np.int64)
        control_qx = np.full((5, 111), np.nan, dtype=np.float64)
        for index, (member, table_id, reviewed_sheet) in enumerate(_CONTROL_MEMBERS):
            counts, ratios = _parse_control_workbook(
                archive.read(member),
                member=member,
                table_id=table_id,
                reviewed_sheet=reviewed_sheet,
            )
            controls[index] = counts
            control_qx[index] = ratios
        _validate_ine_workbook(
            archive.read("tablas_ine/tablas_ine-hombres.xlsx"),
            expected_sex="HOMBRES",
        )
        _validate_ine_workbook(
            archive.read("tablas_ine/tablas_ine-mujeres.xlsx"),
            expected_sex="MUJERES",
        )
        public = _aggregate_public_extract(
            archive,
            expected_row_count=reviewed.expected_public_row_count,
        )

    rows: list[dict[str, object]] = []
    for table_index, table_id in enumerate(reviewed.table_order):
        for age in range(reviewed.age_start, reviewed.age_end + 1):
            public_exposed = int(public[table_index, age, 0])
            public_deaths = int(public[table_index, age, 1])
            workbook_exposed = int(controls[table_index, age, 0])
            workbook_deaths = int(controls[table_index, age, 1])
            delta_exposed = int(deviations[table_index, age, 0])
            delta_deaths = int(deviations[table_index, age, 1])
            exposed_ok = public_exposed - workbook_exposed == delta_exposed
            deaths_ok = public_deaths - workbook_deaths == delta_deaths
            if not exposed_ok:
                raise DataContractError(
                    f"TM-2020 {table_id} age {age} exposed reconciliation failed."
                )
            if not deaths_ok:
                raise DataContractError(
                    f"TM-2020 {table_id} age {age} deaths reconciliation failed."
                )
            if public_exposed < 0 or public_deaths < 0 or public_deaths > public_exposed:
                raise DataContractError(f"TM-2020 {table_id} age {age} public counts are invalid.")
            public_qx = None if public_exposed == 0 else public_deaths / public_exposed
            workbook_qx = None if workbook_exposed == 0 else workbook_deaths / workbook_exposed
            if workbook_exposed > 0 and not math.isclose(
                cast(float, workbook_qx),
                control_qx[table_index, age],
                rel_tol=1e-12,
                abs_tol=1e-15,
            ):
                raise DataContractError(f"TM-2020 {table_id} age {age} workbook qx is invalid.")
            rows.append(
                {
                    "table_id": table_id,
                    "age": age,
                    "public_exposed": public_exposed,
                    "public_deaths": public_deaths,
                    "public_qx": public_qx,
                    "workbook_exposed": workbook_exposed,
                    "workbook_deaths": workbook_deaths,
                    "workbook_qx": workbook_qx,
                    "published_delta_exposed": delta_exposed,
                    "published_delta_deaths": delta_deaths,
                    "exposed_relation_ok": exposed_ok,
                    "deaths_relation_ok": deaths_ok,
                    "qx_reconciled": exposed_ok and deaths_ok,
                }
            )
    frame = pd.DataFrame(rows, columns=_AUDIT_COLUMNS)
    return AnonymizedPublicExtractReconciliation._create(frame)


def parse_cmf_tm2020_construction(
    raw: BinaryIO,
    deviation_raw: bytes,
    contract: CmfMortalityConstructionContract,
) -> AnonymizedPublicExtractReconciliation:
    """Stream public rows without retaining individual data on failure."""

    error_message: str | None = None
    try:
        return _parse_cmf_tm2020_construction(
            raw,
            deviation_raw,
            contract,
        )
    except DataContractError as error:
        error_message = str(error)
    except Exception:
        error_message = "TM-2020 construction release violates its closed parser contract."

    # Raise after the exception handler has exited. This intentionally severs
    # PyArrow/openpyxl exception contexts that may contain source row text.
    del raw, deviation_raw, contract
    raise DataContractError(error_message)
